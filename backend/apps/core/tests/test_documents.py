import base64
import json
import secrets
import threading
import time
import uuid
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from allauth.mfa.totp.internal.auth import TOTP, generate_totp_secret
from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
from defusedxml import ElementTree
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError, close_old_connections, connection, transaction
from django.db.models.deletion import ProtectedError
from django.test import Client
from django.test.utils import override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document as WordDocument

import apps.core.document_exports as document_exports
from apps.accounts.bootstrap import bootstrap_owner
from apps.accounts.models import (
    BuiltInRole,
    CustomRole,
    CustomRolePermission,
    CustomRoleScope,
    ScopedRoleAssignment,
    TenantMembership,
    User,
)
from apps.core.approved_egress import ApprovedTarget
from apps.core.diagram_exports import DiagramExportArtifact, DiagramRenderError, DiagramSource
from apps.core.document_sources import apply_remote_observation, fetch_remote_document, html_to_markdown
from apps.core.documents import PlacementConflict, restructure_document, update_shared_block
from apps.core.models import (
    AuditEvent,
    Block,
    BlockKind,
    BlockRevision,
    Document,
    DocumentationListingReference,
    DocumentAttachment,
    DocumentPlacement,
    DocumentPublication,
    DocumentPublicationArtifact,
    DocumentPublicationControlEvent,
    DocumentRemoteSource,
    DocumentTemplateEnrollment,
    DocumentTopicType,
    Entity,
    EntityLink,
    EntityVisibility,
    InstallationState,
    Organization,
    OrganizationClassification,
    PublicationArtifactKind,
)
from apps.core.publications import canonical_json, publish_document, snapshot_payload, verify_publication
from apps.core.workspaces import resolve_msp_workspace


class UnavailableStorageProvider:
    provider_id = "django-default"

    def quarantine(self, **_kwargs):  # type: ignore[no-untyped-def]
        raise OSError("unavailable")

    def promote(self, **_kwargs):  # type: ignore[no-untyped-def]
        raise OSError("unavailable")

    def read(self, **_kwargs):  # type: ignore[no-untyped-def]
        raise OSError("unavailable")

    def delete(self, **_kwargs):  # type: ignore[no-untyped-def]
        return None


@pytest.fixture
def installation(db):
    InstallationState.objects.get_or_create(pk=InstallationState.SINGLETON_ID)
    result = bootstrap_owner(
        tenant_name="Example MSP",
        owner_email="documents-owner@example.com",
        owner_display_name="Primary Owner",
        password=f"{secrets.token_urlsafe(24)}Aa7!",
    )
    TOTP.activate(result.owner, generate_totp_secret())
    return result


@pytest.fixture
def owner_client(installation):
    client = Client()
    client.force_login(installation.owner)
    return client


@pytest.fixture
def approver_client(installation):
    approver = User.objects.create_user(
        email="documents-approver@example.invalid",
        password=f"{secrets.token_urlsafe(24)}Aa7!",
        display_name="Publication Approver",
    )
    TenantMembership.objects.create(
        tenant=installation.tenant,
        user=approver,
        role=BuiltInRole.ADMINISTRATOR,
    )
    TOTP.activate(approver, generate_totp_secret())
    client = Client()
    client.force_login(approver)
    return client


def organization(tenant, name):  # type: ignore[no-untyped-def]
    entity = Entity.objects.create_owned(tenant=tenant, entity_type="organization", display_name=name)
    record = Organization.objects.create(tenant=tenant, entity=entity, access_mode="all_authorized")
    OrganizationClassification.objects.create(tenant=tenant, organization=record, kind="client")
    return record


@pytest.mark.django_db
def test_document_persists_from_msp_and_client_browser_routes(owner_client, installation):
    client_org = organization(installation.tenant, "Acme Dental")
    msp_url = reverse("msp-document-list-create")
    org_url = reverse("organization-document-list-create", kwargs={"organization_entity_id": client_org.entity_id})

    msp_created = owner_client.post(
        msp_url,
        {"title": "Firewall standard", "markdown": "# Firewall\n\nUse **MFA**."},
        content_type="application/json",
    )
    org_created = owner_client.post(
        org_url, {"title": "Acme onboarding", "markdown": "- Call the client"}, content_type="application/json"
    )

    assert msp_created.status_code == 201
    assert org_created.status_code == 201
    assert msp_created.json()["block_id"]
    assert msp_created.json()["placements"][0]["block_kind"] == BlockKind.RICH_TEXT
    assert msp_created.json()["placements"][0]["resolved_markdown"] == "# Firewall\n\nUse **MFA**."
    assert msp_created.json()["placements"][0]["resolved_html"] == (
        "<h1>Firewall</h1>\n<p>Use <strong>MFA</strong>.</p>\n"
    )
    assert org_created.json()["owner_organization_id"] == str(client_org.entity_id)
    assert Document.objects.count() == 2
    assert Block.objects.count() == 2
    assert DocumentPlacement.objects.values_list("position", flat=True).order_by("position").first() == 0
    detail = reverse("msp-document-detail", kwargs={"document_entity_id": msp_created.json()["id"]})
    updated = owner_client.put(
        detail,
        {
            "title": "Firewall baseline",
            "markdown": "Updated from the browser.",
            "base_revision_id": msp_created.json()["current_revision_id"],
        },
        content_type="application/json",
    )
    assert updated.status_code == 200
    assert updated.json()["title"] == "Firewall baseline"
    assert updated.json()["markdown"] == "Updated from the browser."
    assert updated.json()["revision_number"] == 2
    revisions = list(
        BlockRevision.objects.filter(block__entity_id=msp_created.json()["block_id"]).order_by("revision_number")
    )
    assert [item.markdown for item in revisions] == ["# Firewall\n\nUse **MFA**.", "Updated from the browser."]
    assert revisions[1].parent == revisions[0]
    assert revisions[1].checksum == sha256(b"Updated from the browser.").hexdigest()


@pytest.mark.django_db
def test_structured_topic_conversion_is_explicit_versioned_and_preflighted(owner_client, installation):
    created_response = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Router replacement", "markdown": "Replace the failed router.", "topic_type": "procedure"},
        content_type="application/json",
    )
    assert created_response.status_code == 201
    created = created_response.json()
    assert created["topic_type"] == "procedure"
    assert created["topic_schema_version"] == 1
    assert "## Purpose" in created["markdown"]
    assert "tekdocs:section" not in created["markdown"]
    original_revision = created["current_revision_id"]

    preflight_response = owner_client.get(
        reverse("msp-document-preflight", kwargs={"document_entity_id": created["id"]}),
        {"audience": "msp_internal"},
    )
    assert preflight_response.status_code == 200
    assert preflight_response.json()["valid"] is True
    assert preflight_response.json()["version"] == "tekdocs-preflight/v1"
    assert not [item for item in preflight_response.json()["findings"] if item["severity"] == "blocker"]

    conversion_url = reverse("msp-document-topic-conversion", kwargs={"document_entity_id": created["id"]})
    preview_response = owner_client.post(
        conversion_url,
        {"topic_type": "troubleshooting", "base_revision_id": original_revision, "apply": False},
        content_type="application/json",
    )
    assert preview_response.status_code == 200
    assert preview_response.json()["topic_type"] == "troubleshooting"
    assert "## Issue" in preview_response.json()["converted_markdown"]
    assert Document.objects.get(entity_id=created["id"]).topic_type == "procedure"

    applied_response = owner_client.post(
        conversion_url,
        {"topic_type": "troubleshooting", "base_revision_id": original_revision, "apply": True},
        content_type="application/json",
    )
    assert applied_response.status_code == 200
    assert applied_response.json()["topic_type"] == "troubleshooting"
    assert applied_response.json()["current_revision_id"] != original_revision
    historical = BlockRevision.objects.get(id=original_revision)
    assert historical.topic_type == "procedure"
    assert "## Purpose" in historical.markdown
    stale_preview = owner_client.post(
        conversion_url,
        {"topic_type": "reference", "base_revision_id": original_revision, "apply": False},
        content_type="application/json",
    )
    assert stale_preview.status_code == 409
    assert stale_preview.json()["code"] == "revision_conflict"


@pytest.mark.django_db
def test_policy_starter_markdown_is_accepted_without_duplicate_structure(owner_client, installation):
    from apps.core.topic_schemas import seed_markdown

    starter = seed_markdown(DocumentTopicType.POLICY)
    response = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Acceptable use", "markdown": starter, "topic_type": "policy", "category": "policy"},
        content_type="application/json",
    )
    assert response.status_code == 201
    assert response.json()["markdown"] == starter
    assert response.json()["markdown"].count("## Purpose") == 1
    assert "tekdocs:section" not in response.json()["markdown"]


@pytest.mark.django_db
def test_remote_document_source_retains_reviewable_change_before_explicit_apply(
    owner_client, installation, monkeypatch
):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Published setup", "markdown": "Original\n"},
        content_type="application/json",
    )
    document_id = created.json()["id"]
    source_response = owner_client.put(
        reverse("msp-document-remote-source", kwargs={"document_entity_id": document_id}),
        {"url": "https://docs.example.invalid/setup", "source_kind": "html", "check_interval_minutes": 60},
        content_type="application/json",
    )
    assert source_response.status_code == 201

    class Response:
        status = 200
        headers = {"Content-Type": "text/html", "ETag": "private-upstream-value"}

        def read(self, _size):  # type: ignore[no-untyped-def]
            return b"<article><h1>Setup</h1><p>Use MFA.</p><script>ignore()</script></article>"

        def close(self):
            return None

    class Pool:
        def urlopen(self, *_args, **_kwargs):  # type: ignore[no-untyped-def]
            return Response()

        def close(self):
            return None

    monkeypatch.setattr(
        "apps.core.document_sources.resolve_public_https_target",
        lambda *_args, **_kwargs: ApprovedTarget(
            "https://docs.example.invalid/setup", "docs.example.invalid", "203.0.113.10", "/setup"
        ),
    )
    monkeypatch.setattr("apps.core.document_sources.pinned_https_pool", lambda *_args, **_kwargs: Pool())
    checked = owner_client.post(
        reverse("msp-document-remote-observations", kwargs={"document_entity_id": document_id}),
        {},
        content_type="application/json",
    )
    assert checked.status_code == 201
    assert checked.json()["state"] == "changed"
    assert "Use MFA." in checked.json()["canonical_markdown"]
    assert "ignore" not in checked.json()["canonical_markdown"]
    assert "private-upstream-value" not in str(checked.json())
    assert created.json()["markdown"] == "Original\n"

    applied = owner_client.post(
        reverse(
            "msp-document-remote-observation-apply",
            kwargs={"document_entity_id": document_id, "observation_id": checked.json()["id"]},
        ),
        {},
        content_type="application/json",
    )
    assert applied.status_code == 200
    reopened = owner_client.get(reverse("msp-document-detail", kwargs={"document_entity_id": document_id}))
    assert reopened.json()["markdown"].startswith("# Setup")


@pytest.mark.django_db
def test_remote_document_source_rejects_non_https_and_retains_value_free_fetch_failures(
    owner_client, installation, monkeypatch
):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Remote source", "markdown": "Safe"},
        content_type="application/json",
    )
    source_url = reverse("msp-document-remote-source", kwargs={"document_entity_id": created.json()["id"]})
    rejected = owner_client.put(source_url, {"url": "http://127.0.0.1/private"}, content_type="application/json")
    assert rejected.status_code == 400
    accepted = owner_client.put(
        source_url, {"url": "https://public.example.invalid/doc"}, content_type="application/json"
    )
    assert accepted.status_code == 201
    monkeypatch.setattr(
        "apps.core.document_sources.resolve_public_https_target",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(ValidationError("Document source destination denied")),
    )
    observation = fetch_remote_document(DocumentRemoteSource.objects.get(id=accepted.json()["id"]))
    assert observation.state == "failed"
    assert observation.error_code == "remote_connection_failed"
    assert observation.canonical_markdown == ""
    with pytest.raises(PlacementConflict):
        apply_remote_observation(observation=observation, actor_id=installation.owner.id)


def test_remote_html_conversion_is_semantic_and_drops_executable_content():
    markdown = html_to_markdown("<h2>Network</h2><ul><li>LAN</li></ul><style>hidden</style>")
    assert markdown == "## Network\n- LAN\n"


@pytest.mark.django_db
def test_document_categories_templates_and_filters_persist(owner_client, installation):
    collection = reverse("msp-document-list-create")
    created = owner_client.post(
        collection,
        {"title": "Access policy", "markdown": "Policy body", "category": "policy", "is_template": True},
        content_type="application/json",
    )
    assert created.status_code == 201
    assert created.json()["category"] == "policy"
    assert created.json()["is_template"] is True

    templates = owner_client.get(collection, {"category": "policy", "template": "templates", "q": "access"})
    documents = owner_client.get(collection, {"template": "documents"})
    assert [item["title"] for item in templates.json()["results"]] == ["Access policy"]
    assert documents.json()["count"] == 0


@pytest.mark.django_db
def test_static_publication_freezes_dependencies_and_verifies_after_source_changes(
    owner_client, installation, tmp_path
):
    client_org = organization(installation.tenant, "Static Client")
    collection = reverse("organization-document-list-create", kwargs={"organization_entity_id": client_org.entity_id})
    with override_settings(MEDIA_ROOT=tmp_path):
        created = owner_client.post(
            collection,
            {"title": "Access standard", "markdown": "Initial", "category": "policy"},
            content_type="application/json",
        ).json()
        private_asset = Entity.objects.create_owned(
            tenant=installation.tenant,
            organization=client_org,
            entity_type="client_asset",
            display_name="Private publishing note",
            visibility=EntityVisibility.MSP_PRIVATE,
        )
        EntityLink.objects.create(
            tenant=installation.tenant,
            source_id=created["id"],
            target=private_asset,
            link_type="references",
        )
        attachment = owner_client.post(
            reverse(
                "organization-document-attachment-list-create",
                kwargs={"organization_entity_id": client_org.entity_id, "document_entity_id": created["id"]},
            ),
            {"file": SimpleUploadedFile("evidence.txt", b"retained evidence", content_type="text/html")},
        ).json()
        source_markdown = (
            "# Access standard\n\n"
            f"[Client](tekdocs://entity/{client_org.entity_id})\n\n"
            f"[Evidence](tekdocs://attachment/{attachment['id']})"
        )
        detail_url = reverse(
            "organization-document-detail",
            kwargs={"organization_entity_id": client_org.entity_id, "document_entity_id": created["id"]},
        )
        updated = owner_client.put(
            detail_url,
            {
                "title": "Access standard",
                "markdown": source_markdown,
                "base_revision_id": created["current_revision_id"],
                "category": "policy",
            },
            content_type="application/json",
        ).json()
        publication_url = reverse(
            "organization-document-publication-list-create",
            kwargs={"organization_entity_id": client_org.entity_id, "document_entity_id": created["id"]},
        )
        response = owner_client.post(
            publication_url,
            {"reason": "Approved access policy", "audience": "client_visible", "retention": "permanent"},
            content_type="application/json",
        )

        assert response.status_code == 201
        publication_payload = response.json()
        assert publication_payload["canonical_markdown"] == source_markdown + "\n"
        assert publication_payload["verification"] == {
            "valid": True,
            "digest_valid": True,
            "signature_valid": True,
            "key_fingerprint_valid": True,
            "trusted_key": True,
        }
        assert publication_payload["manifest"]["format"] == "tekdocs-static-publication/v4"
        assert publication_payload["manifest"]["placements"][0]["revision_id"] == updated["current_revision_id"]
        assert publication_payload["manifest"]["entities"][0]["display_name"] == "Static Client"
        assert publication_payload["manifest"]["attachments"][0]["checksum"] == attachment["checksum"]
        assert publication_payload["manifest"]["relationship_graph"]["family"] == "document"
        assert publication_payload["manifest"]["relationship_graph"]["root_entity_id"] == created["id"]
        assert {item["label"] for item in publication_payload["manifest"]["relationship_graph"]["nodes"]} == {
            "Access standard"
        }
        assert publication_payload["reason"] == "Approved access policy"
        assert publication_payload["audience"] == "client_visible"
        assert publication_payload["lifecycle_state"] == "pending_approval"
        assert publication_payload["audience_projections"][1] == {
            "audience": "client_portal",
            "available": False,
            "state": "pending_approval",
        }
        assert len(publication_payload["artifacts"]) == 2
        assert {item["kind"] for item in publication_payload["manifest"]["artifacts"]} == {"pdf", "attachment"}
        assert "Access standard" in publication_payload["sanitized_html"]
        assert 'href="tekdocs:' not in publication_payload["sanitized_html"]

        source_changed = owner_client.put(
            detail_url,
            {
                "title": "Changed source",
                "markdown": "Replacement content",
                "base_revision_id": updated["current_revision_id"],
                "category": "guide",
            },
            content_type="application/json",
        )
        assert source_changed.status_code == 200
        publication = DocumentPublication.objects.get(entity_id=publication_payload["id"])
        assert publication.title == "Access standard"
        assert publication.canonical_markdown == source_markdown + "\n"
        assert verify_publication(publication)["valid"] is True
        publication.canonical_markdown += "tampered"
        assert verify_publication(publication) == {
            "valid": False,
            "digest_valid": False,
            "signature_valid": False,
            "key_fingerprint_valid": True,
            "trusted_key": True,
        }
        publication.refresh_from_db()
        attacker_key = Ed25519PrivateKey.generate()
        attacker_public = attacker_key.public_key().public_bytes(
            encoding=serialization.Encoding.Raw,
            format=serialization.PublicFormat.Raw,
        )
        digest = sha256(
            snapshot_payload(
                manifest=publication.manifest,
                markdown=publication.canonical_markdown,
                sanitized_html=publication.sanitized_html,
            )
        ).digest()
        publication.public_key = base64.urlsafe_b64encode(attacker_public).decode("ascii")
        publication.key_fingerprint = sha256(attacker_public).hexdigest()
        publication.signature = base64.urlsafe_b64encode(attacker_key.sign(digest)).decode("ascii")
        forged = verify_publication(publication)
        assert forged["signature_valid"] is True
        assert forged["key_fingerprint_valid"] is True
        assert forged["trusted_key"] is False
        assert forged["valid"] is False
        publication.refresh_from_db()

        publication_detail = reverse(
            "organization-document-publication-detail",
            kwargs={
                "organization_entity_id": client_org.entity_id,
                "document_entity_id": created["id"],
                "publication_entity_id": publication.entity_id,
            },
        )
        assert owner_client.get(publication_detail).json()["content_digest"] == publication.content_digest
        markdown_download = owner_client.get(publication_detail + "/markdown")
        manifest_download = owner_client.get(publication_detail + "/manifest")
        assert markdown_download.content == publication.canonical_markdown.encode()
        assert manifest_download.content == canonical_json(publication.manifest) + b"\n"
        assert markdown_download["Cache-Control"] == "private, no-store"
        assert manifest_download["X-Content-Type-Options"] == "nosniff"
        pdf_artifact = next(item for item in publication_payload["artifacts"] if item["kind"] == "pdf")
        retained_attachment = next(item for item in publication_payload["artifacts"] if item["kind"] == "attachment")
        artifact_route = "organization-document-publication-artifact-download"
        artifact_kwargs = {
            "organization_entity_id": client_org.entity_id,
            "document_entity_id": created["id"],
            "publication_entity_id": publication.entity_id,
        }
        pdf_download = owner_client.get(
            reverse(artifact_route, kwargs={**artifact_kwargs, "artifact_entity_id": pdf_artifact["id"]})
        )
        retained_download = owner_client.get(
            reverse(artifact_route, kwargs={**artifact_kwargs, "artifact_entity_id": retained_attachment["id"]})
        )
        assert pdf_download.content.startswith(b"%PDF-")
        assert retained_download.content == b"retained evidence"
        assert pdf_download["Cache-Control"] == "private, no-store"
        assert retained_download["X-Content-Type-Options"] == "nosniff"
        assert all(
            artifact.original_filename not in artifact.file.name
            for artifact in DocumentPublicationArtifact.objects.all()
        )


@pytest.mark.django_db
def test_static_correction_retention_and_audience_rules_are_append_only(
    owner_client, approver_client, installation, tmp_path
):
    client_org = organization(installation.tenant, "Lifecycle Client")
    created = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": client_org.entity_id}),
        {"title": "Retention policy", "markdown": "Approved content"},
        content_type="application/json",
    ).json()
    publication_url = reverse(
        "organization-document-publication-list-create",
        kwargs={"organization_entity_id": client_org.entity_id, "document_entity_id": created["id"]},
    )
    with override_settings(MEDIA_ROOT=tmp_path):
        first = owner_client.post(
            publication_url,
            {
                "reason": "Annual approval",
                "audience": "client_visible",
                "retention": "review_on",
                "retention_review_on": timezone.localdate().isoformat(),
            },
            content_type="application/json",
        )
        assert first.status_code == 201
        assert first.json()["lifecycle_state"] == "pending_approval"
        approval_url = reverse(
            "organization-document-publication-approve",
            kwargs={
                "organization_entity_id": client_org.entity_id,
                "document_entity_id": created["id"],
                "publication_entity_id": first.json()["id"],
            },
        )
        self_approval = owner_client.post(
            approval_url,
            {"reason": "Attempted self approval"},
            content_type="application/json",
        )
        assert self_approval.status_code == 409
        approved_first = approver_client.post(
            approval_url,
            {"reason": "Independent policy approval"},
            content_type="application/json",
        )
        assert approved_first.status_code == 200
        assert approved_first.json()["lifecycle_state"] == "review_due"
        assert approved_first.json()["audience_projections"][1]["available"] is True
        corrected = owner_client.post(
            publication_url,
            {
                "reason": "Corrected an outdated support contact",
                "audience": "client_visible",
                "retention": "permanent",
                "supersedes_id": first.json()["id"],
            },
            content_type="application/json",
        )
        assert corrected.status_code == 201
        assert corrected.json()["supersedes_id"] == first.json()["id"]
        prior = DocumentPublication.objects.get(entity_id=first.json()["id"])
        assert prior.lifecycle_state == "review_due"
        corrected_approval = approver_client.post(
            reverse(
                "organization-document-publication-approve",
                kwargs={
                    "organization_entity_id": client_org.entity_id,
                    "document_entity_id": created["id"],
                    "publication_entity_id": corrected.json()["id"],
                },
            ),
            {"reason": "Correction independently approved"},
            content_type="application/json",
        )
        assert corrected_approval.status_code == 200
        assert corrected_approval.json()["lifecycle_state"] == "published"
        prior.refresh_from_db()
        assert prior.lifecycle_state == "superseded"
        duplicate = owner_client.post(
            publication_url,
            {
                "reason": "Second replacement",
                "audience": "client_visible",
                "retention": "permanent",
                "supersedes_id": first.json()["id"],
            },
            content_type="application/json",
        )
        assert duplicate.status_code == 409
        assert DocumentPublication.objects.count() == 2

        withdrawal = approver_client.post(
            reverse(
                "organization-document-publication-withdraw",
                kwargs={
                    "organization_entity_id": client_org.entity_id,
                    "document_entity_id": created["id"],
                    "publication_entity_id": corrected.json()["id"],
                },
            ),
            {"reason": "Policy withdrawn pending replacement"},
            content_type="application/json",
        )
        assert withdrawal.status_code == 200
        assert withdrawal.json()["lifecycle_state"] == "withdrawn"
        assert withdrawal.json()["audience_projections"][1] == {
            "audience": "client_portal",
            "available": False,
            "state": "withdrawn",
        }

    msp_document = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Internal only", "markdown": "MSP"},
        content_type="application/json",
    ).json()
    invalid_audience = owner_client.post(
        reverse("msp-document-publication-list-create", kwargs={"document_entity_id": msp_document["id"]}),
        {"reason": "Wrong audience", "audience": "client_visible", "retention": "permanent"},
        content_type="application/json",
    )
    assert invalid_audience.status_code == 400
    missing_reason = owner_client.post(
        reverse("msp-document-publication-list-create", kwargs={"document_entity_id": msp_document["id"]}),
        {"audience": "msp_internal", "retention": "permanent"},
        content_type="application/json",
    )
    assert missing_reason.status_code == 400


@pytest.mark.django_db
def test_static_publication_fails_closed_for_unavailable_dependencies(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {
            "title": "Broken dependency",
            "markdown": f"[Missing](tekdocs://entity/{uuid.uuid4()})",
        },
        content_type="application/json",
    ).json()
    response = owner_client.post(
        reverse("msp-document-publication-list-create", kwargs={"document_entity_id": created["id"]}),
        {"reason": "Dependency test", "audience": "msp_internal", "retention": "permanent"},
        content_type="application/json",
    )
    assert response.status_code == 409
    assert response.json()["code"] == "publication_conflict"
    assert DocumentPublication.objects.count() == 0
    assert Entity.objects.filter(entity_type="document_publication").count() == 0


@pytest.mark.django_db
def test_static_publication_retains_exact_diagram_artifacts(owner_client, installation, tmp_path, monkeypatch):
    markdown = """```mermaid
flowchart LR
  accTitle: Service request
  accDescr: Browser reaches the API.
  Browser --> API
```
"""
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Diagram publication", "markdown": markdown},
        content_type="application/json",
    ).json()
    source = DiagramSource(
        index=1,
        source="flowchart LR\n  accTitle: Service request\n  accDescr: Browser reaches the API.\n  Browser --> API",
        source_checksum=sha256(
            b"flowchart LR\n  accTitle: Service request\n  accDescr: Browser reaches the API.\n  Browser --> API"
        ).hexdigest(),
        title="Service request",
        description="Browser reaches the API.",
    )
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    diagram = DiagramExportArtifact(
        source=source,
        state="rendered",
        renderer_version="renderer-test",
        svg=b"<svg><title>Service request</title><desc>Browser reaches the API.</desc></svg>",
        png=png,
    )
    monkeypatch.setattr("apps.core.publications.render_diagram_exports", lambda *_args, **_kwargs: (diagram,))
    monkeypatch.setattr("apps.core.preflight.render_diagram_exports", lambda *_args, **_kwargs: (diagram,))

    with override_settings(MEDIA_ROOT=tmp_path):
        response = owner_client.post(
            reverse("msp-document-publication-list-create", kwargs={"document_entity_id": created["id"]}),
            {"reason": "Diagram retention", "audience": "msp_internal", "retention": "permanent"},
            content_type="application/json",
        )

    assert response.status_code == 201
    publication = DocumentPublication.objects.get(entity_id=response.json()["id"])
    assert [item["state"] for item in publication.manifest["diagrams"]] == ["rendered"]
    assert publication.manifest["diagrams"][0]["source_checksum"] == source.source_checksum
    assert "data:image/svg+xml;base64," in publication.sanitized_html
    assert "Accessible diagram source" in publication.sanitized_html
    artifacts = publication.artifacts.order_by("kind")
    assert {item.kind for item in artifacts} == {
        PublicationArtifactKind.PDF,
        PublicationArtifactKind.DIAGRAM_SVG,
        PublicationArtifactKind.DIAGRAM_PNG,
    }
    assert all(item.checksum and item.size > 0 for item in artifacts)


@pytest.mark.django_db
def test_static_publication_does_not_promote_when_diagram_rendering_fails(owner_client, installation, monkeypatch):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Rejected diagram", "markdown": "```mermaid\nflowchart LR\nA-->B\n```\n"},
        content_type="application/json",
    ).json()

    def fail_render(*_args, **_kwargs):  # type: ignore[no-untyped-def]
        raise DiagramRenderError("failed")

    monkeypatch.setattr("apps.core.publications.render_diagram_exports", fail_render)
    response = owner_client.post(
        reverse("msp-document-publication-list-create", kwargs={"document_entity_id": created["id"]}),
        {"reason": "Must fail", "audience": "msp_internal", "retention": "permanent"},
        content_type="application/json",
    )
    assert response.status_code == 409
    assert response.json()["code"] == "publication_conflict"
    assert DocumentPublication.objects.count() == 0
    assert DocumentPublicationArtifact.objects.count() == 0


@pytest.mark.django_db
def test_static_publication_cleans_retained_files_when_artifact_persistence_fails(
    owner_client, installation, tmp_path, monkeypatch
):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Storage rollback", "markdown": "Retain nothing after failure."},
        content_type="application/json",
    ).json()
    document = Document.objects.get(entity_id=created["id"])

    def reject_artifact_save(*args, **kwargs):  # type: ignore[no-untyped-def]
        raise OSError("simulated retained storage failure")

    monkeypatch.setattr(DocumentPublicationArtifact, "save", reject_artifact_save)
    with override_settings(MEDIA_ROOT=tmp_path), pytest.raises(OSError, match="simulated"):
        publish_document(
            workspace=resolve_msp_workspace(installation.owner),
            document=document,
            actor_id=installation.owner.pk,
            reason="Storage rollback test",
            audience="msp_internal",
            retention="permanent",
            retention_review_on=None,
        )

    assert DocumentPublication.objects.count() == 0
    assert DocumentPublicationArtifact.objects.count() == 0
    assert not [path for path in tmp_path.rglob("*") if path.is_file()]


@pytest.mark.django_db
def test_static_publication_requires_owning_workspace_and_hides_cross_client_ids(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    beta = organization(installation.tenant, "Beta")
    shared = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Shared standard", "markdown": "MSP-owned"},
        content_type="application/json",
    ).json()
    owner_client.post(
        reverse("msp-document-reference-list-create", kwargs={"document_entity_id": shared["id"]}),
        {"organization_id": str(acme.entity_id)},
        content_type="application/json",
    )
    referenced_publish = owner_client.post(
        reverse(
            "organization-document-publication-list-create",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": shared["id"]},
        ),
        {"reason": "Invalid workspace", "audience": "msp_internal", "retention": "permanent"},
        content_type="application/json",
    )
    assert referenced_publish.status_code == 404

    acme_document = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
        {"title": "Acme-only", "markdown": "Private"},
        content_type="application/json",
    ).json()
    published = owner_client.post(
        reverse(
            "organization-document-publication-list-create",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": acme_document["id"]},
        ),
        {"reason": "Client record", "audience": "msp_internal", "retention": "permanent"},
        content_type="application/json",
    ).json()
    cross_client = owner_client.get(
        reverse(
            "organization-document-publication-detail",
            kwargs={
                "organization_entity_id": beta.entity_id,
                "document_entity_id": acme_document["id"],
                "publication_entity_id": published["id"],
            },
        )
    )
    assert cross_client.status_code == 404
    pdf_artifact = next(artifact for artifact in published["artifacts"] if artifact["kind"] == "pdf")
    cross_client_artifact = owner_client.get(
        reverse(
            "organization-document-publication-artifact-download",
            kwargs={
                "organization_entity_id": beta.entity_id,
                "document_entity_id": acme_document["id"],
                "publication_entity_id": published["id"],
                "artifact_entity_id": pdf_artifact["id"],
            },
        )
    )
    assert cross_client_artifact.status_code == 404


@pytest.mark.django_db(transaction=True)
def test_static_publication_is_append_only_in_django_and_postgresql(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Immutable publication", "markdown": "Frozen"},
        content_type="application/json",
    ).json()
    owner_client.post(
        reverse("msp-document-publication-list-create", kwargs={"document_entity_id": created["id"]}),
        {"reason": "Immutability test", "audience": "msp_internal", "retention": "permanent"},
        content_type="application/json",
    )
    publication = DocumentPublication.objects.get()
    artifact = DocumentPublicationArtifact.objects.get(kind="pdf")
    control_event = DocumentPublicationControlEvent.objects.get(action="approved")
    publication.title = "Mutated"
    with pytest.raises(ValidationError, match="append-only"):
        publication.save()
    with pytest.raises(ValidationError, match="append-only"):
        publication.delete()
    artifact.original_filename = "mutated.pdf"
    with pytest.raises(ValidationError, match="append-only"):
        artifact.save()
    with pytest.raises(ValidationError, match="append-only"):
        artifact.delete()
    control_event.reason = "Mutated"
    with pytest.raises(ValidationError, match="append-only"):
        control_event.save()
    with pytest.raises(ValidationError, match="append-only"):
        control_event.delete()
    malformed = DocumentPublication(
        tenant=publication.tenant,
        document=publication.document,
        entity=Entity.objects.create_owned(
            tenant=publication.tenant,
            entity_type="document_publication",
            display_name="Malformed publication",
        ),
        title="Malformed",
        category="general",
        reason="Malformed test",
        audience="msp_internal",
        retention="permanent",
        manifest={},
        content_digest="a" * 64,
        signature="fixture",
        public_key="fixture",
        key_fingerprint="b" * 64,
        published_at=publication.published_at,
    )
    with pytest.raises(ValidationError, match="manifest format"):
        malformed.full_clean()
    if connection.vendor == "postgresql":
        with pytest.raises(DatabaseError, match="append-only"), transaction.atomic():
            DocumentPublication.objects.filter(pk=publication.pk).update(title="Database mutation")
        with pytest.raises(ProtectedError):
            DocumentPublication.objects.filter(pk=publication.pk).delete()
        with pytest.raises(DatabaseError, match="append-only"), transaction.atomic(), connection.cursor() as cursor:
            cursor.execute("DELETE FROM core_documentpublication WHERE id = %s", [publication.pk])
        with pytest.raises(DatabaseError, match="append-only"), transaction.atomic():
            DocumentPublicationArtifact.objects.filter(pk=artifact.pk).update(original_filename="database.pdf")
        with pytest.raises(DatabaseError, match="append-only"), transaction.atomic(), connection.cursor() as cursor:
            cursor.execute("DELETE FROM core_documentpublicationartifact WHERE id = %s", [artifact.pk])
        with pytest.raises(DatabaseError, match="append-only"), transaction.atomic():
            DocumentPublicationControlEvent.objects.filter(pk=control_event.pk).update(reason="Database mutation")
        with pytest.raises(DatabaseError, match="append-only"), transaction.atomic(), connection.cursor() as cursor:
            cursor.execute("DELETE FROM core_documentpublicationcontrolevent WHERE id = %s", [control_event.pk])
        with pytest.raises(DatabaseError, match="requires an actor"), transaction.atomic():
            DocumentPublicationControlEvent.objects.create(
                tenant=publication.tenant,
                publication=publication,
                action="withdrawn",
                reason="Forged anonymous withdrawal",
            )
        with pytest.raises(DatabaseError, match="publication"), transaction.atomic():
            DocumentPublication.objects.bulk_create([malformed])
        malformed_v3_id = uuid.uuid4()
        malformed_v3_entity = Entity.objects.create_owned(
            tenant=publication.tenant,
            entity_type="document_publication",
            display_name="Malformed v3 publication",
        )
        malformed_v3_manifest = json.loads(json.dumps(publication.manifest))
        malformed_v3_manifest["publication_id"] = str(malformed_v3_id)
        malformed_v3_manifest["publication_entity_id"] = str(malformed_v3_entity.id)
        malformed_v3_manifest["key_resolutions"] = [{"kind": "field"}]
        malformed_v3 = DocumentPublication(
            id=malformed_v3_id,
            tenant=publication.tenant,
            organization=publication.organization,
            document=publication.document,
            entity=malformed_v3_entity,
            title=publication.title,
            category=publication.category,
            reason=publication.reason,
            audience=publication.audience,
            retention=publication.retention,
            retention_review_on=publication.retention_review_on,
            canonical_markdown=publication.canonical_markdown,
            sanitized_html=publication.sanitized_html,
            manifest=malformed_v3_manifest,
            content_digest=publication.content_digest,
            signature=publication.signature,
            signature_algorithm=publication.signature_algorithm,
            public_key=publication.public_key,
            key_fingerprint=publication.key_fingerprint,
            published_by=publication.published_by,
            published_at=publication.published_at,
        )
        with pytest.raises(DatabaseError, match="key resolution"), transaction.atomic():
            DocumentPublication.objects.bulk_create([malformed_v3])


@pytest.mark.django_db
def test_template_instantiation_copies_owned_attachments_and_rewrites_markdown(owner_client, installation, tmp_path):
    with override_settings(MEDIA_ROOT=tmp_path):
        template = owner_client.post(
            reverse("msp-document-list-create"),
            {"title": "Router template", "markdown": "Start", "category": "guide", "is_template": True},
            content_type="application/json",
        ).json()
        upload = owner_client.post(
            reverse("msp-document-attachment-list-create", kwargs={"document_entity_id": template["id"]}),
            {"file": SimpleUploadedFile("router.pdf", b"%PDF-1.4\nTekDocs fixture\n%%EOF", content_type="text/html")},
        )
        assert upload.status_code == 201
        attachment_id = upload.json()["id"]
        detail = reverse("msp-document-detail", kwargs={"document_entity_id": template["id"]})
        linked = owner_client.put(
            detail,
            {
                "title": "Router template",
                "markdown": f"[authored label](tekdocs://attachment/{attachment_id})",
                "base_revision_id": template["current_revision_id"],
                "category": "guide",
                "is_template": True,
            },
            content_type="application/json",
        ).json()
        created = owner_client.post(
            reverse("msp-document-template-instantiate"),
            {"source_document_id": template["id"], "title": "Branch router", "category": "procedure"},
            content_type="application/json",
        )
        assert created.status_code == 201
        payload = created.json()
        assert payload["category"] == "procedure"
        assert payload["is_template"] is False
        assert payload["attachment_count"] == 1
        copied_id = payload["attachments"][0]["id"]
        assert copied_id != attachment_id
        assert f"tekdocs://attachment/{copied_id}" in payload["markdown"]
        assert attachment_id not in payload["markdown"]
        assert payload["revision_number"] == 1
        assert linked["revision_number"] == 2
        assert DocumentAttachment.objects.count() == 2


@pytest.mark.django_db
def test_client_template_creates_independent_copy_only_within_its_workspace(owner_client, installation):
    acme = organization(installation.tenant, "Acme local templates")
    beta = organization(installation.tenant, "Beta local templates")
    template = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
        {
            "title": "Acme incident template",
            "markdown": "Client-owned response steps.",
            "category": "procedure",
            "is_template": True,
        },
        content_type="application/json",
    ).json()

    created = owner_client.post(
        reverse("organization-document-template-instantiate", kwargs={"organization_entity_id": acme.entity_id}),
        {
            "source_document_id": template["id"],
            "title": "Acme incident response",
            "category": "procedure",
        },
        content_type="application/json",
    )
    assert created.status_code == 201
    assert created.json()["markdown"] == "Client-owned response steps."
    assert created.json()["is_template"] is False
    assert created.json()["template_enrollment_id"] is None

    cross_client = owner_client.post(
        reverse("organization-document-template-instantiate", kwargs={"organization_entity_id": beta.entity_id}),
        {
            "source_document_id": template["id"],
            "title": "Leaked incident response",
            "category": "procedure",
        },
        content_type="application/json",
    )
    assert cross_client.status_code == 404


@pytest.mark.django_db
def test_managed_attachment_download_is_private_and_scope_bound(owner_client, installation, tmp_path):
    acme = organization(installation.tenant, "Acme attachments")
    beta = organization(installation.tenant, "Beta attachments")
    with override_settings(MEDIA_ROOT=tmp_path):
        created = owner_client.post(
            reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
            {"title": "Private file", "markdown": "", "category": "reference"},
            content_type="application/json",
        ).json()
        upload_url = reverse(
            "organization-document-attachment-list-create",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": created["id"]},
        )
        rejected = owner_client.post(
            upload_url,
            {"file": SimpleUploadedFile("payload.svg", b"<svg><script>alert(1)</script></svg>")},
        )
        assert rejected.status_code == 400
        uploaded = owner_client.post(
            upload_url,
            {"file": SimpleUploadedFile("notes.txt", b"safe UTF-8 notes", content_type="text/html")},
        )
        assert uploaded.status_code == 201
        attachment = DocumentAttachment.objects.get(entity_id=uploaded.json()["id"])
        assert "notes.txt" not in attachment.file.name
        assert attachment.media_type == "text/plain"
        assert attachment.scan_status == "clean"
        assert attachment.scan_engine == "tekdocs-strict-content/1"
        assert attachment.storage_provider == "django-default"
        download = owner_client.get(
            reverse(
                "organization-document-attachment-download",
                kwargs={
                    "organization_entity_id": acme.entity_id,
                    "document_entity_id": created["id"],
                    "attachment_entity_id": attachment.entity_id,
                },
            )
        )
        assert b"".join(download.streaming_content) == b"safe UTF-8 notes"
        assert download["Content-Type"] == "application/octet-stream"
        assert download["Cache-Control"] == "private, no-store"
        assert download["X-Content-Type-Options"] == "nosniff"
        assert download["Accept-Ranges"] == "bytes"
        partial = owner_client.get(
            reverse(
                "organization-document-attachment-download",
                kwargs={
                    "organization_entity_id": acme.entity_id,
                    "document_entity_id": created["id"],
                    "attachment_entity_id": attachment.entity_id,
                },
            ),
            HTTP_RANGE="bytes=5-8",
        )
        assert partial.status_code == 206
        assert partial.content == b"UTF-"
        assert partial["Content-Range"] == "bytes 5-8/16"
        suffix = owner_client.get(
            reverse(
                "organization-document-attachment-download",
                kwargs={
                    "organization_entity_id": acme.entity_id,
                    "document_entity_id": created["id"],
                    "attachment_entity_id": attachment.entity_id,
                },
            ),
            HTTP_RANGE="bytes=-5",
        )
        assert suffix.status_code == 206
        assert suffix.content == b"notes"
        unsatisfiable = owner_client.get(
            reverse(
                "organization-document-attachment-download",
                kwargs={
                    "organization_entity_id": acme.entity_id,
                    "document_entity_id": created["id"],
                    "attachment_entity_id": attachment.entity_id,
                },
            ),
            HTTP_RANGE="bytes=999-1000",
        )
        assert unsatisfiable.status_code == 416
        assert unsatisfiable["Content-Range"] == "bytes */16"
        download_events = AuditEvent.objects.filter(
            action="document.attachment.downloaded", entity_id=attachment.entity_id
        )
        assert sorted(event.metadata["partial"] for event in download_events) == [False, True, True]
        assert all(set(event.metadata) == {"partial", "purpose"} for event in download_events)
        stored_name = attachment.file.name
        attachment.file.storage.delete(stored_name)
        assert attachment.file.storage.save(stored_name, ContentFile(b"tampered bytes")) == stored_name
        assert (
            owner_client.get(
                reverse(
                    "organization-document-attachment-download",
                    kwargs={
                        "organization_entity_id": acme.entity_id,
                        "document_entity_id": created["id"],
                        "attachment_entity_id": attachment.entity_id,
                    },
                )
            ).status_code
            == 400
        )
        sibling = reverse(
            "organization-document-attachment-download",
            kwargs={
                "organization_entity_id": beta.entity_id,
                "document_entity_id": created["id"],
                "attachment_entity_id": attachment.entity_id,
            },
        )
        assert owner_client.get(sibling).status_code == 404


@pytest.mark.django_db
def test_file_backed_document_creation_and_primary_file_replacement_are_versioned(owner_client, installation, tmp_path):
    acme = organization(installation.tenant, "Acme file documents")
    beta = organization(installation.tenant, "Beta file documents")
    create_url = reverse(
        "organization-document-file-backed-create",
        kwargs={"organization_entity_id": acme.entity_id},
    )
    with override_settings(MEDIA_ROOT=tmp_path):
        created = owner_client.post(
            create_url,
            {
                "title": "Firewall datasheet",
                "notes": "## Deployment notes\r\n\r\nRetain the approved model.",
                "category": "reference",
                "file": SimpleUploadedFile(
                    "firewall-v1.pdf",
                    b"%PDF-1.4\nTekDocs primary fixture v1\n%%EOF",
                    content_type="text/html",
                ),
            },
        )
        assert created.status_code == 201
        payload = created.json()
        assert payload["markdown"] == "## Deployment notes\n\nRetain the approved model."
        assert payload["attachments"] == []
        assert payload["attachment_count"] == 0
        assert payload["primary_file"]["filename"] == "firewall-v1.pdf"
        assert payload["primary_file"]["media_type"] == "application/pdf"
        assert payload["primary_file"]["version_number"] == 1
        assert payload["primary_file"]["is_current"] is True
        assert payload["primary_file"]["replaces_id"] is None

        replace_url = reverse(
            "organization-document-primary-file",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": payload["id"]},
        )
        replaced = owner_client.post(
            replace_url,
            {
                "file": SimpleUploadedFile(
                    "firewall-v2.pdf",
                    b"%PDF-1.4\nTekDocs primary fixture v2\n%%EOF",
                    content_type="application/pdf",
                )
            },
        )
        assert replaced.status_code == 201
        assert replaced.json()["version_number"] == 2
        assert replaced.json()["replaces_id"] == payload["primary_file"]["id"]

        refreshed = owner_client.get(
            reverse(
                "organization-document-detail",
                kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": payload["id"]},
            )
        ).json()
        assert refreshed["primary_file"]["filename"] == "firewall-v2.pdf"
        assert [item["version_number"] for item in refreshed["primary_file_versions"]] == [2, 1]
        assert [item["is_current"] for item in refreshed["primary_file_versions"]] == [True, False]

        current = DocumentAttachment.objects.get(entity_id=refreshed["primary_file"]["id"])
        current.original_filename = "mutated.pdf"
        with pytest.raises(ValidationError, match="immutable"):
            current.save()
        with pytest.raises(ValidationError, match="immutable"):
            current.delete()
        if connection.vendor == "postgresql":
            with pytest.raises(DatabaseError, match="immutable"), transaction.atomic():
                DocumentAttachment.objects.filter(pk=current.pk).update(original_filename="database-mutated.pdf")
            with pytest.raises(DatabaseError, match="immutable"), transaction.atomic(), connection.cursor() as cursor:
                cursor.execute("DELETE FROM core_documentattachment WHERE id = %s", [current.pk])

        for version, expected in zip(
            refreshed["primary_file_versions"],
            (b"%PDF-1.4\nTekDocs primary fixture v2\n%%EOF", b"%PDF-1.4\nTekDocs primary fixture v1\n%%EOF"),
            strict=True,
        ):
            download = owner_client.get(
                reverse(
                    "organization-document-attachment-download",
                    kwargs={
                        "organization_entity_id": acme.entity_id,
                        "document_entity_id": payload["id"],
                        "attachment_entity_id": version["id"],
                    },
                )
            )
            assert b"".join(download.streaming_content) == expected
            assert download["Cache-Control"] == "private, no-store"
            assert download["Content-Type"] == "application/octet-stream"

        assert (
            owner_client.post(
                reverse(
                    "organization-document-primary-file",
                    kwargs={"organization_entity_id": beta.entity_id, "document_entity_id": payload["id"]},
                ),
                {"file": SimpleUploadedFile("hidden.pdf", b"%PDF-1.4\nhidden\n%%EOF")},
            ).status_code
            == 404
        )
        assert (
            owner_client.delete(
                reverse(
                    "organization-document-attachment-detail",
                    kwargs={
                        "organization_entity_id": acme.entity_id,
                        "document_entity_id": payload["id"],
                        "attachment_entity_id": payload["primary_file"]["id"],
                    },
                )
            ).status_code
            == 400
        )

    assert AuditEvent.objects.filter(action="document.primary_file.created").count() == 2
    assert all(
        set(event.metadata) <= {"document_id", "version_number"}
        for event in AuditEvent.objects.filter(action="document.primary_file.created")
    )


@pytest.mark.django_db
def test_rejected_file_backed_document_creation_leaves_no_document(owner_client, installation, tmp_path):
    before_documents = Document.objects.count()
    before_entities = Entity.objects.count()
    with override_settings(MEDIA_ROOT=tmp_path):
        rejected = owner_client.post(
            reverse("msp-document-file-backed-create"),
            {
                "title": "Rejected file",
                "notes": "This must roll back.",
                "category": "reference",
                "file": SimpleUploadedFile("payload.pdf", b"MZ executable content", content_type="application/pdf"),
            },
        )
    assert rejected.status_code == 400
    assert Document.objects.count() == before_documents
    assert Entity.objects.count() == before_entities


@pytest.mark.django_db
def test_markdown_import_and_resolved_export_use_portable_bytes(owner_client, installation):
    imported = owner_client.post(
        reverse("msp-document-import"),
        {
            "file": SimpleUploadedFile("switch-guide.md", "# Switch\n\nCafé\n".encode()),
            "title": "Imported switch guide",
            "category": "guide",
            "is_template": "false",
        },
    )
    assert imported.status_code == 201
    payload = imported.json()
    assert payload["markdown"] == "# Switch"
    assert payload["resolved_markdown"] == "# Switch\n\nCafé\n"
    assert [placement["block_kind"] for placement in payload["placements"]] == ["heading", "rich_text"]
    exported = owner_client.get(reverse("msp-document-export", kwargs={"document_entity_id": payload["id"]}))
    assert exported.content == "# Switch\n\nCafé\n".encode()
    assert exported["Content-Type"].startswith("text/markdown")
    assert "imported-switch-guide-editable.md" in exported["Content-Disposition"]
    assert exported["X-TekDocs-Export-Class"] == "editable_revision_snapshot"
    assert exported["Cache-Control"] == "private, no-store"


@pytest.mark.django_db
def test_msp_document_reference_cross_lists_without_copying_or_lateral_disclosure(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    beta = organization(installation.tenant, "Beta")
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Shared response guide", "markdown": "One canonical block."},
        content_type="application/json",
    ).json()
    references_url = reverse("msp-document-reference-list-create", kwargs={"document_entity_id": created["id"]})
    added = owner_client.post(references_url, {"organization_id": str(acme.entity_id)}, content_type="application/json")
    assert added.status_code == 201
    assert Document.objects.count() == 1
    assert Block.objects.count() == 1
    assert DocumentationListingReference.objects.count() == 1

    acme_url = reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id})
    beta_url = reverse("organization-document-list-create", kwargs={"organization_entity_id": beta.entity_id})
    assert owner_client.get(acme_url).json()["results"][0]["is_reference"] is True
    assert owner_client.get(acme_url).json()["results"][0]["title"] == "Shared response guide"
    assert owner_client.get(beta_url).json()["count"] == 0

    removed = owner_client.delete(
        reverse(
            "msp-document-reference-detail",
            kwargs={"document_entity_id": created["id"], "reference_id": added.json()["id"]},
        )
    )
    assert removed.status_code == 204
    assert owner_client.get(acme_url).json()["count"] == 0


@pytest.mark.django_db
def test_document_route_scope_returns_not_found_for_another_client(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    beta = organization(installation.tenant, "Beta")
    acme_url = reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id})
    created = owner_client.post(
        acme_url, {"title": "Acme private", "markdown": "Never disclose laterally."}, content_type="application/json"
    ).json()
    beta_detail = reverse(
        "organization-document-detail",
        kwargs={"organization_entity_id": beta.entity_id, "document_entity_id": created["id"]},
    )
    assert owner_client.get(beta_detail).status_code == 404
    assert (
        owner_client.put(
            beta_detail,
            {"title": "Changed", "markdown": "Bad", "base_revision_id": created["current_revision_id"]},
            content_type="application/json",
        ).status_code
        == 404
    )


@pytest.mark.django_db
def test_stale_document_update_returns_diff_without_overwriting_draft(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Runbook", "markdown": "line one\nline two\n"},
        content_type="application/json",
    ).json()
    detail = reverse("msp-document-detail", kwargs={"document_entity_id": created["id"]})
    first = owner_client.put(
        detail,
        {"title": "Runbook", "markdown": "line one\nserver edit\n", "base_revision_id": created["current_revision_id"]},
        content_type="application/json",
    ).json()
    stale = owner_client.put(
        detail,
        {"title": "Runbook", "markdown": "line one\nmy draft\n", "base_revision_id": created["current_revision_id"]},
        content_type="application/json",
    )
    assert stale.status_code == 409
    assert stale.json()["code"] == "revision_conflict"
    assert stale.json()["current_revision"]["id"] == first["current_revision_id"]
    assert "+server edit" in stale.json()["diff"]
    assert owner_client.get(detail).json()["markdown"] == "line one\nserver edit\n"
    assert BlockRevision.objects.count() == 2


@pytest.mark.django_db
def test_revision_history_is_scoped_and_noop_content_does_not_add_revision(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    beta = organization(installation.tenant, "Beta")
    collection = reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id})
    created = owner_client.post(
        collection, {"title": "Access guide", "markdown": "First"}, content_type="application/json"
    ).json()
    detail = reverse(
        "organization-document-detail",
        kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": created["id"]},
    )
    renamed = owner_client.put(
        detail,
        {"title": "Access procedure", "markdown": "First", "base_revision_id": created["current_revision_id"]},
        content_type="application/json",
    ).json()
    assert renamed["revision_number"] == 1
    history_url = reverse(
        "organization-document-revision-list",
        kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": created["id"]},
    )
    history = owner_client.get(history_url)
    assert history.status_code == 200
    assert history.json()["count"] == 1
    assert history.json()["page"] == 1
    assert history.json()["page_size"] == 50
    assert history.json()["has_more"] is False
    assert len(history.json()["results"]) == 1
    revision_url = reverse(
        "organization-document-revision-detail",
        kwargs={
            "organization_entity_id": acme.entity_id,
            "document_entity_id": created["id"],
            "revision_id": created["current_revision_id"],
        },
    )
    assert owner_client.get(revision_url).json()["markdown"] == "First"
    beta_history = reverse(
        "organization-document-revision-list",
        kwargs={"organization_entity_id": beta.entity_id, "document_entity_id": created["id"]},
    )
    assert owner_client.get(beta_history).status_code == 404


@pytest.mark.django_db
def test_revision_history_is_paginated_and_rejects_unbounded_pages(owner_client, installation):
    collection = reverse("msp-document-list-create")
    record = owner_client.post(
        collection, {"title": "Long-lived runbook", "markdown": "Revision 1"}, content_type="application/json"
    ).json()
    detail = reverse("msp-document-detail", kwargs={"document_entity_id": record["id"]})
    for revision_number in range(2, 53):
        record = owner_client.put(
            detail,
            {
                "title": record["title"],
                "markdown": f"Revision {revision_number}",
                "base_revision_id": record["current_revision_id"],
            },
            content_type="application/json",
        ).json()
    history_url = reverse("msp-document-revision-list", kwargs={"document_entity_id": record["id"]})
    first = owner_client.get(history_url).json()
    second = owner_client.get(history_url, {"page": 2, "page_size": 50}).json()
    assert (first["count"], first["page"], first["has_more"], len(first["results"])) == (52, 1, True, 50)
    assert first["results"][0]["revision_number"] == 52
    assert [item["revision_number"] for item in second["results"]] == [2, 1]
    assert second["has_more"] is False
    assert owner_client.get(history_url, {"page_size": 101}).status_code == 400


@pytest.mark.django_db(transaction=True)
def test_block_revisions_are_append_only_in_application_and_postgresql(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Immutable runbook", "markdown": "Original"},
        content_type="application/json",
    ).json()
    revision = BlockRevision.objects.get(id=created["current_revision_id"])
    revision.markdown = "Mutated"
    with pytest.raises(ValidationError):
        revision.save()
    if connection.vendor == "postgresql":
        with pytest.raises(DatabaseError), transaction.atomic(), connection.cursor() as cursor:
            cursor.execute("UPDATE core_blockrevision SET markdown = %s WHERE id = %s", ["Mutated", revision.id])
    revision.refresh_from_db()
    assert revision.markdown == "Original"


@pytest.mark.django_db
def test_live_and_pinned_transclusions_resolve_deterministically(owner_client, installation):
    collection = reverse("msp-document-list-create")
    source = owner_client.post(
        collection, {"title": "Shared checklist", "markdown": "Shared revision one"}, content_type="application/json"
    ).json()
    live_document = owner_client.post(
        collection, {"title": "Live composition", "markdown": "Live introduction"}, content_type="application/json"
    ).json()
    pinned_document = owner_client.post(
        collection, {"title": "Pinned composition", "markdown": "Pinned introduction"}, content_type="application/json"
    ).json()

    live_added = owner_client.post(
        reverse("msp-document-placement-list-create", kwargs={"document_entity_id": live_document["id"]}),
        {"source_document_id": source["id"], "resolution_mode": "live"},
        content_type="application/json",
    )
    pinned_added = owner_client.post(
        reverse("msp-document-placement-list-create", kwargs={"document_entity_id": pinned_document["id"]}),
        {
            "source_document_id": source["id"],
            "resolution_mode": "pinned",
            "pinned_revision_id": source["current_revision_id"],
        },
        content_type="application/json",
    )
    assert live_added.status_code == 200
    assert pinned_added.status_code == 200
    assert live_added.json()["resolved_markdown"] == "Live introduction\n\nShared revision one\n"
    assert pinned_added.json()["resolved_markdown"] == "Pinned introduction\n\nShared revision one\n"

    owner_client.put(
        reverse("msp-document-detail", kwargs={"document_entity_id": source["id"]}),
        {
            "title": source["title"],
            "markdown": "Shared revision two",
            "base_revision_id": source["current_revision_id"],
        },
        content_type="application/json",
    )
    live_reloaded = owner_client.get(
        reverse("msp-document-detail", kwargs={"document_entity_id": live_document["id"]})
    ).json()
    pinned_reloaded = owner_client.get(
        reverse("msp-document-detail", kwargs={"document_entity_id": pinned_document["id"]})
    ).json()
    assert live_reloaded["resolved_markdown"].endswith("Shared revision two\n")
    assert pinned_reloaded["resolved_markdown"].endswith("Shared revision one\n")
    assert live_reloaded["placements"][1]["resolved_revision_number"] == 2
    assert pinned_reloaded["placements"][1]["resolved_revision_number"] == 1


@pytest.mark.django_db
def test_document_composition_creates_independent_typed_blocks_at_exact_positions(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Block-native guide", "markdown": "Introduction"},
        content_type="application/json",
    ).json()
    placements_url = reverse("msp-document-placement-list-create", kwargs={"document_entity_id": created["id"]})

    code = owner_client.post(
        placements_url,
        {
            "operation": "create_block",
            "block_kind": "code",
            "block_name": "Validation command",
            "markdown": "```shell\nmake check\n```",
        },
        content_type="application/json",
    )
    heading = owner_client.post(
        placements_url,
        {
            "operation": "create_block",
            "block_kind": "heading",
            "block_name": "Preparation heading",
            "markdown": "## Preparation",
            "position": 1,
        },
        content_type="application/json",
    )

    assert code.status_code == 200
    assert heading.status_code == 200
    payload = heading.json()
    assert payload["markdown"] == "Introduction"
    assert payload["resolved_markdown"] == "Introduction\n\n## Preparation\n\n```shell\nmake check\n```\n"
    assert [(item["position"], item["block_kind"]) for item in payload["placements"]] == [
        (0, "rich_text"),
        (1, "heading"),
        (2, "code"),
    ]
    assert payload["placements"][1]["resolved_markdown"] == "## Preparation"
    assert payload["placements"][2]["resolved_markdown"] == "```shell\nmake check\n```"
    blocks = list(Block.objects.filter(tenant=installation.tenant).order_by("created_at"))
    assert [block.kind for block in blocks] == ["rich_text", "code", "heading"]
    assert all(block.entity.entity_type == "document_block" for block in blocks)
    document = Document.objects.get(entity_id=created["id"])
    assert all(block.source_document_id == document.id for block in blocks)
    assert BlockRevision.objects.filter(block__in=blocks, revision_number=1).count() == 3


@pytest.mark.django_db
def test_removing_owned_block_archives_its_identity_without_orphaning_content(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Lifecycle", "markdown": "Primary"},
        content_type="application/json",
    ).json()
    placements_url = reverse("msp-document-placement-list-create", kwargs={"document_entity_id": created["id"]})
    added = owner_client.post(
        placements_url,
        {"operation": "create_block", "block_kind": "rich_text", "markdown": "Temporary"},
        content_type="application/json",
    ).json()
    secondary = added["placements"][1]

    removed = owner_client.delete(
        reverse(
            "msp-document-placement-detail",
            kwargs={"document_entity_id": created["id"], "placement_id": secondary["id"]},
        )
    )

    assert removed.status_code == 200
    block = Block.objects.get(entity_id=secondary["block_id"])
    assert block.archived_at is not None
    assert block.entity.archived_at is not None
    assert not block.placements.exists()


@pytest.mark.django_db
def test_placement_audiences_filter_publications_before_snapshot_provenance(owner_client, installation, tmp_path):
    client_org = organization(installation.tenant, "Audience Client")
    route_kwargs = {"organization_entity_id": client_org.entity_id}
    created = owner_client.post(
        reverse("organization-document-list-create", kwargs=route_kwargs),
        {"title": "Audience runbook", "markdown": "Shared introduction"},
        content_type="application/json",
    ).json()
    placement_route = reverse(
        "organization-document-placement-list-create",
        kwargs={**route_kwargs, "document_entity_id": created["id"]},
    )
    private_entity = Entity.objects.create_owned(
        tenant=installation.tenant,
        organization=client_org,
        entity_type="client_asset",
        display_name="Internal audience asset",
        visibility=EntityVisibility.MSP_PRIVATE,
    )
    with override_settings(MEDIA_ROOT=tmp_path):
        internal_attachment = owner_client.post(
            reverse(
                "organization-document-attachment-list-create",
                kwargs={**route_kwargs, "document_entity_id": created["id"]},
            ),
            {"file": SimpleUploadedFile("internal-evidence.txt", b"INTERNAL-FILE-SENTINEL")},
        ).json()
    internal_markdown = (
        "INTERNAL-SENTINEL\n\n"
        f"[Internal asset](tekdocs://entity/{private_entity.id})\n\n"
        f"[Internal evidence](tekdocs://attachment/{internal_attachment['id']})"
    )
    internal = owner_client.post(
        placement_route,
        {
            "operation": "create_block",
            "block_name": "Operator secret",
            "markdown": internal_markdown,
            "audience_profile": "msp_internal",
        },
        content_type="application/json",
    ).json()["placements"][1]
    composed = owner_client.post(
        placement_route,
        {
            "operation": "create_block",
            "block_name": "Client instructions",
            "markdown": "CLIENT-SENTINEL",
            "audience_profile": "client_visible",
        },
        content_type="application/json",
    ).json()
    client_placement = composed["placements"][2]

    assert [item["audience_profile"] for item in composed["placements"]] == [
        "shared",
        "msp_internal",
        "client_visible",
    ]
    publication_route = reverse(
        "organization-document-publication-list-create",
        kwargs={**route_kwargs, "document_entity_id": created["id"]},
    )
    with override_settings(MEDIA_ROOT=tmp_path):
        client_publication = owner_client.post(
            publication_route,
            {"reason": "Client edition", "audience": "client_visible", "retention": "permanent"},
            content_type="application/json",
        )
        msp_publication = owner_client.post(
            publication_route,
            {"reason": "Operator edition", "audience": "msp_internal", "retention": "permanent"},
            content_type="application/json",
        )

    assert client_publication.status_code == 201
    assert msp_publication.status_code == 201
    client_payload = client_publication.json()
    msp_payload = msp_publication.json()
    assert client_payload["canonical_markdown"] == "Shared introduction\n\nCLIENT-SENTINEL\n"
    assert msp_payload["canonical_markdown"] == f"Shared introduction\n\n{internal_markdown}\n"
    assert [item["id"] for item in client_payload["manifest"]["placements"]] == [
        composed["placements"][0]["id"],
        client_placement["id"],
    ]
    assert [item["id"] for item in msp_payload["manifest"]["placements"]] == [
        composed["placements"][0]["id"],
        internal["id"],
    ]
    assert "INTERNAL-SENTINEL" not in json.dumps(client_payload)
    assert internal["id"] not in json.dumps(client_payload)
    assert str(private_entity.id) not in json.dumps(client_payload)
    assert internal_attachment["id"] not in json.dumps(client_payload)
    assert "CLIENT-SENTINEL" not in json.dumps(msp_payload)
    assert client_placement["id"] not in json.dumps(msp_payload)
    assert msp_payload["manifest"]["entities"][0]["id"] == str(private_entity.id)
    assert msp_payload["manifest"]["attachments"][0]["id"] == internal_attachment["id"]


@pytest.mark.django_db
def test_placement_audience_rejects_widening_nested_or_primary_content(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Audience hierarchy", "markdown": "Primary"},
        content_type="application/json",
    ).json()
    placement_route = reverse("msp-document-placement-list-create", kwargs={"document_entity_id": created["id"]})
    parent = owner_client.post(
        placement_route,
        {
            "operation": "create_block",
            "markdown": "Internal parent",
            "audience_profile": "msp_internal",
        },
        content_type="application/json",
    ).json()["placements"][1]
    widened = owner_client.post(
        placement_route,
        {
            "operation": "create_block",
            "markdown": "Shared child",
            "parent_id": parent["id"],
            "audience_profile": "shared",
        },
        content_type="application/json",
    )
    valid_child = owner_client.post(
        placement_route,
        {
            "operation": "create_block",
            "markdown": "Internal child",
            "parent_id": parent["id"],
            "audience_profile": "msp_internal",
        },
        content_type="application/json",
    ).json()["placements"][-1]
    primary_update = owner_client.patch(
        reverse(
            "msp-document-placement-detail",
            kwargs={"document_entity_id": created["id"], "placement_id": created["placements"][0]["id"]},
        ),
        {"audience_profile": "client_visible"},
        content_type="application/json",
    )

    assert widened.status_code == 409
    assert "cannot widen" in widened.json()["detail"]
    assert primary_update.status_code == 409
    assert "live and shared" in primary_update.json()["detail"]
    with pytest.raises(DatabaseError, match="child placement cannot widen parent audience"), transaction.atomic():
        DocumentPlacement.objects.filter(id=valid_child["id"]).update(audience_profile="shared")
    with pytest.raises(DatabaseError, match="primary document placement must remain shared"), transaction.atomic():
        DocumentPlacement.objects.filter(id=created["placements"][0]["id"]).update(audience_profile="client_visible")


@pytest.mark.django_db
def test_excluded_placement_keys_are_not_resolved_for_other_audiences(owner_client, installation, tmp_path):
    client_org = organization(installation.tenant, "Audience Key Client")
    route_kwargs = {"organization_entity_id": client_org.entity_id}
    created = owner_client.post(
        reverse("organization-document-list-create", kwargs=route_kwargs),
        {"title": "Audience key boundary", "markdown": "Shared content"},
        content_type="application/json",
    ).json()
    owner_client.post(
        reverse(
            "organization-document-placement-list-create",
            kwargs={**route_kwargs, "document_entity_id": created["id"]},
        ),
        {
            "operation": "create_block",
            "markdown": "Internal <tekdocs://key/missing.serial_number>",
            "audience_profile": "msp_internal",
        },
        content_type="application/json",
    )
    publication_route = reverse(
        "organization-document-publication-list-create",
        kwargs={**route_kwargs, "document_entity_id": created["id"]},
    )
    with override_settings(MEDIA_ROOT=tmp_path):
        client_publication = owner_client.post(
            publication_route,
            {"reason": "Client key boundary", "audience": "client_visible", "retention": "permanent"},
            content_type="application/json",
        )
        msp_publication = owner_client.post(
            publication_route,
            {"reason": "MSP key boundary", "audience": "msp_internal", "retention": "permanent"},
            content_type="application/json",
        )

    assert client_publication.status_code == 201
    assert client_publication.json()["canonical_markdown"] == "Shared content\n"
    assert client_publication.json()["manifest"]["key_resolutions"] == []
    assert msp_publication.status_code == 409
    assert DocumentPublication.objects.count() == 1


@pytest.mark.django_db
def test_archiving_document_archives_every_owned_block(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Archive composition", "markdown": "Primary"},
        content_type="application/json",
    ).json()
    owner_client.post(
        reverse("msp-document-placement-list-create", kwargs={"document_entity_id": created["id"]}),
        {"operation": "create_block", "block_kind": "code", "markdown": "```shell\ntrue\n```"},
        content_type="application/json",
    )
    document = Document.objects.get(entity_id=created["id"])
    assert document.owned_blocks.count() == 2

    archived = owner_client.delete(reverse("msp-document-detail", kwargs={"document_entity_id": created["id"]}))

    assert archived.status_code == 204
    assert not document.owned_blocks.filter(archived_at__isnull=True).exists()
    assert not Entity.objects.filter(block_record__source_document=document, archived_at__isnull=True).exists()


@pytest.mark.django_db
def test_new_document_block_contract_rejects_source_and_pinned_semantics(owner_client, installation):
    source = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Source", "markdown": "Source body"},
        content_type="application/json",
    ).json()
    destination = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Destination", "markdown": "Destination body"},
        content_type="application/json",
    ).json()
    placements_url = reverse(
        "msp-document-placement-list-create",
        kwargs={"document_entity_id": destination["id"]},
    )

    mixed = owner_client.post(
        placements_url,
        {
            "operation": "create_block",
            "source_document_id": source["id"],
            "markdown": "Not allowed",
        },
        content_type="application/json",
    )
    pinned = owner_client.post(
        placements_url,
        {
            "operation": "create_block",
            "markdown": "Not allowed",
            "resolution_mode": "pinned",
            "pinned_revision_id": source["current_revision_id"],
        },
        content_type="application/json",
    )
    invalid_position = owner_client.post(
        placements_url,
        {"operation": "create_block", "markdown": "Not allowed", "position": 99},
        content_type="application/json",
    )

    assert mixed.status_code == 400
    assert pinned.status_code == 400
    assert invalid_position.status_code == 409
    assert Document.objects.get(entity_id=destination["id"]).placements.count() == 1


@pytest.mark.django_db(transaction=True)
def test_database_rejects_unknown_document_block_kind(owner_client, installation):
    if connection.vendor != "postgresql":
        pytest.skip("Block-kind database constraint requires PostgreSQL")
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Kind guard", "markdown": "Guarded"},
        content_type="application/json",
    ).json()
    with pytest.raises(DatabaseError), transaction.atomic():
        Block.objects.filter(entity_id=created["block_id"]).update(kind="script")


@pytest.mark.django_db
def test_nested_transclusion_is_depth_first_and_rejects_ancestor_cycle(owner_client, installation):
    collection = reverse("msp-document-list-create")
    destination = owner_client.post(
        collection, {"title": "Composition", "markdown": "A"}, content_type="application/json"
    ).json()
    source_b = owner_client.post(collection, {"title": "B", "markdown": "B"}, content_type="application/json").json()
    source_c = owner_client.post(collection, {"title": "C", "markdown": "C"}, content_type="application/json").json()
    placements_url = reverse("msp-document-placement-list-create", kwargs={"document_entity_id": destination["id"]})
    added_b = owner_client.post(
        placements_url,
        {"source_document_id": source_b["id"], "resolution_mode": "live"},
        content_type="application/json",
    ).json()
    b_placement = added_b["placements"][1]
    added_c = owner_client.post(
        placements_url,
        {
            "source_document_id": source_c["id"],
            "resolution_mode": "live",
            "parent_id": b_placement["id"],
        },
        content_type="application/json",
    )
    assert added_c.status_code == 200
    assert added_c.json()["resolved_markdown"] == "A\n\nB\n\nC\n"
    assert [item["depth"] for item in added_c.json()["placements"]] == [0, 0, 1]

    cycle = owner_client.post(
        placements_url,
        {
            "source_document_id": source_b["id"],
            "resolution_mode": "live",
            "parent_id": b_placement["id"],
        },
        content_type="application/json",
    )
    assert cycle.status_code == 409
    assert cycle.json()["code"] == "placement_conflict"
    assert "Circular" in cycle.json()["detail"]


@pytest.mark.django_db
def test_transclusion_source_must_be_visible_in_destination_workspace(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    beta = organization(installation.tenant, "Beta")
    acme_collection = reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id})
    beta_collection = reverse("organization-document-list-create", kwargs={"organization_entity_id": beta.entity_id})
    destination = owner_client.post(
        acme_collection, {"title": "Acme composition", "markdown": "Acme"}, content_type="application/json"
    ).json()
    sibling_source = owner_client.post(
        beta_collection, {"title": "Beta private", "markdown": "Beta"}, content_type="application/json"
    ).json()
    response = owner_client.post(
        reverse(
            "organization-document-placement-list-create",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": destination["id"]},
        ),
        {"source_document_id": sibling_source["id"], "resolution_mode": "live"},
        content_type="application/json",
    )
    assert response.status_code == 404


@pytest.mark.django_db
def test_client_block_library_requires_explicit_msp_opt_in_and_never_exposes_sibling_clients(
    owner_client, installation
):
    acme = organization(installation.tenant, "Acme Library")
    beta = organization(installation.tenant, "Beta Library")
    msp_collection = reverse("msp-document-list-create")
    public_source = owner_client.post(
        msp_collection,
        {
            "title": "Published MSP standard",
            "markdown": "Public primary guidance",
            "library_visible": True,
        },
        content_type="application/json",
    ).json()
    private_source = owner_client.post(
        msp_collection,
        {"title": "Private MSP notes", "markdown": "Never disclose"},
        content_type="application/json",
    ).json()
    public_secondary = owner_client.post(
        reverse("msp-document-placement-list-create", kwargs={"document_entity_id": public_source["id"]}),
        {
            "operation": "create_block",
            "block_name": "Printer isolation rationale",
            "markdown": "Printers belong on IoT.",
            "library_visible": True,
        },
        content_type="application/json",
    ).json()["placements"][1]
    private_secondary = owner_client.post(
        reverse("msp-document-placement-list-create", kwargs={"document_entity_id": public_source["id"]}),
        {"operation": "create_block", "block_name": "Internal exception", "markdown": "Private"},
        content_type="application/json",
    ).json()["placements"][2]
    beta_source = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": beta.entity_id}),
        {"title": "Beta-only runbook", "markdown": "Sibling secret", "library_visible": True},
        content_type="application/json",
    ).json()
    destination = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
        {"title": "Acme runbook", "markdown": "Local introduction"},
        content_type="application/json",
    ).json()

    library_url = reverse("organization-document-block-library", kwargs={"organization_entity_id": acme.entity_id})
    library = owner_client.get(library_url)
    assert library.status_code == 200
    ids = {item["id"] for item in library.json()["results"]}
    assert public_source["block_id"] in ids
    assert public_secondary["block_id"] in ids
    assert destination["block_id"] in ids
    assert private_source["block_id"] not in ids
    assert private_secondary["block_id"] not in ids
    assert beta_source["block_id"] not in ids
    assert "Never disclose" not in library.content.decode()
    assert "Sibling secret" not in library.content.decode()

    placement_url = reverse(
        "organization-document-placement-list-create",
        kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": destination["id"]},
    )
    reused = owner_client.post(
        placement_url,
        {"operation": "reuse_block", "source_block_id": public_secondary["block_id"], "resolution_mode": "live"},
        content_type="application/json",
    )
    assert reused.status_code == 200
    assert reused.json()["placements"][-1]["resolved_markdown"] == "Printers belong on IoT."
    visibility_update = owner_client.put(
        reverse("msp-document-detail", kwargs={"document_entity_id": public_source["id"]}),
        {
            "title": public_source["title"],
            "markdown": public_source["markdown"],
            "base_revision_id": public_source["current_revision_id"],
            "category": public_source["category"],
            "is_template": False,
            "library_visible": False,
        },
        content_type="application/json",
    )
    assert visibility_update.status_code == 409
    assert "Detach client block reuse" in visibility_update.json()["detail"]
    assert (
        owner_client.post(
            placement_url,
            {"operation": "reuse_block", "source_block_id": private_secondary["block_id"], "resolution_mode": "live"},
            content_type="application/json",
        ).status_code
        == 404
    )
    assert (
        owner_client.post(
            placement_url,
            {"operation": "reuse_block", "source_block_id": beta_source["block_id"], "resolution_mode": "live"},
            content_type="application/json",
        ).status_code
        == 404
    )


@pytest.mark.django_db
def test_versioned_client_template_rollout_preserves_copy_and_updates_only_explicit_safe_changes(
    owner_client, installation
):
    acme = organization(installation.tenant, "Acme Template")
    beta = organization(installation.tenant, "Beta Template")
    template = owner_client.post(
        reverse("msp-document-list-create"),
        {
            "title": "Client network baseline",
            "markdown": "Client-specific introduction",
            "category": "guide",
            "is_template": True,
            "library_visible": True,
        },
        content_type="application/json",
    ).json()
    second = owner_client.post(
        reverse("msp-document-placement-list-create", kwargs={"document_entity_id": template["id"]}),
        {
            "operation": "create_block",
            "block_name": "Printer rationale",
            "markdown": "Printers belong on IoT.",
            "library_visible": True,
        },
        content_type="application/json",
    ).json()["placements"][1]

    library = owner_client.get(
        reverse("organization-document-template-library", kwargs={"organization_entity_id": acme.entity_id})
    )
    assert library.status_code == 200
    assert [record["id"] for record in library.json()["results"]] == [template["id"]]

    created = owner_client.post(
        reverse("organization-document-template-instantiate", kwargs={"organization_entity_id": acme.entity_id}),
        {
            "source_document_id": template["id"],
            "title": "Acme network baseline",
            "category": "guide",
            "placement_rules": {second["block_id"]: "live"},
        },
        content_type="application/json",
    )
    assert created.status_code == 201
    assert created.json()["placements"][0]["block_id"] != template["block_id"]
    assert created.json()["placements"][1]["block_id"] == second["block_id"]
    enrollment = DocumentTemplateEnrollment.objects.get(destination_document__entity_id=created.json()["id"])

    source_placement = DocumentPlacement.objects.get(id=second["id"])
    update_shared_block(
        placement=source_placement,
        actor_id=installation.owner.id,
        markdown="Printers and media devices belong on IoT.",
        base_revision_id=uuid.UUID(second["resolved_revision_id"]),
    )
    added = owner_client.post(
        reverse("msp-document-placement-list-create", kwargs={"document_entity_id": template["id"]}),
        {
            "operation": "create_block",
            "block_name": "Firewall zones",
            "markdown": "Document the allowed flows.",
            "library_visible": True,
        },
        content_type="application/json",
    ).json()["placements"][2]

    preview_url = reverse(
        "organization-document-template-rollout-preview", kwargs={"organization_entity_id": acme.entity_id}
    )
    preview = owner_client.post(preview_url, {"enrollment_id": str(enrollment.id)}, content_type="application/json")
    assert preview.status_code == 200
    assert preview.json()["up_to_date"] is False
    assert [item["source_block_id"] for item in preview.json()["added"]] == [added["block_id"]]
    assert preview.json()["changed"][0]["mode"] == "live"
    assert preview.json()["conflicts"] == []

    applied = owner_client.post(
        reverse("organization-document-template-rollout-apply", kwargs={"organization_entity_id": acme.entity_id}),
        {
            "enrollment_id": str(enrollment.id),
            "expected_applied_revision_id": preview.json()["applied_revision_id"],
            "placement_rules": {added["block_id"]: "copy"},
        },
        content_type="application/json",
    )
    assert applied.status_code == 200
    enrollment.refresh_from_db()
    assert enrollment.applied_revision_id == uuid.UUID(applied.json()["applied_revision_id"])
    destination = owner_client.get(
        reverse(
            "organization-document-detail",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": created.json()["id"]},
        )
    ).json()
    assert destination["placement_count"] == 3
    assert "Printers and media devices belong on IoT." in destination["resolved_markdown"]
    assert (
        owner_client.post(
            reverse(
                "organization-document-template-rollout-preview", kwargs={"organization_entity_id": beta.entity_id}
            ),
            {"enrollment_id": str(enrollment.id)},
            content_type="application/json",
        ).status_code
        == 404
    )


@pytest.mark.django_db
def test_referenced_msp_block_can_be_transcluded_but_reference_cannot_be_revoked_while_used(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    source = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "MSP standard", "markdown": "Shared standard"},
        content_type="application/json",
    ).json()
    reference = owner_client.post(
        reverse("msp-document-reference-list-create", kwargs={"document_entity_id": source["id"]}),
        {"organization_id": str(acme.entity_id)},
        content_type="application/json",
    ).json()
    destination = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
        {"title": "Client runbook", "markdown": "Client introduction"},
        content_type="application/json",
    ).json()
    added = owner_client.post(
        reverse(
            "organization-document-placement-list-create",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": destination["id"]},
        ),
        {"source_document_id": source["id"], "resolution_mode": "live"},
        content_type="application/json",
    )
    assert added.status_code == 200
    assert added.json()["resolved_markdown"] == "Client introduction\n\nShared standard\n"

    revoked = owner_client.delete(
        reverse(
            "msp-document-reference-detail",
            kwargs={"document_entity_id": source["id"], "reference_id": reference["id"]},
        )
    )
    assert revoked.status_code == 409
    assert revoked.json()["code"] == "placement_conflict"
    archived = owner_client.delete(reverse("msp-document-detail", kwargs={"document_entity_id": source["id"]}))
    assert archived.status_code == 409
    assert Document.objects.get(entity_id=source["id"]).archived_at is None


@pytest.mark.django_db(transaction=True)
def test_postgresql_rejects_raw_placement_cycle_and_cross_client_block(owner_client, installation):
    if connection.vendor != "postgresql":
        pytest.skip("Placement database guards require PostgreSQL")
    acme = organization(installation.tenant, "Acme")
    beta = organization(installation.tenant, "Beta")
    destination = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
        {"title": "Acme composition", "markdown": "A"},
        content_type="application/json",
    ).json()
    beta_source = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": beta.entity_id}),
        {"title": "Beta source", "markdown": "B"},
        content_type="application/json",
    ).json()
    destination_record = Document.objects.get(entity_id=destination["id"])
    root = destination_record.placements.get(parent__isnull=True, position=0)
    beta_block = Block.objects.get(entity_id=beta_source["block_id"])
    with pytest.raises(DatabaseError), transaction.atomic():
        DocumentPlacement.objects.create(
            tenant=installation.tenant,
            organization=acme,
            document=destination_record,
            block=beta_block,
            position=1,
        )

    source = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
        {"title": "Acme source", "markdown": "B"},
        content_type="application/json",
    ).json()
    added = owner_client.post(
        reverse(
            "organization-document-placement-list-create",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": destination["id"]},
        ),
        {"source_document_id": source["id"], "resolution_mode": "live", "parent_id": str(root.id)},
        content_type="application/json",
    ).json()
    child = DocumentPlacement.objects.get(id=added["placements"][1]["id"])
    with pytest.raises(DatabaseError), transaction.atomic():
        DocumentPlacement.objects.filter(id=root.id).update(parent=child)
    with pytest.raises(DatabaseError), transaction.atomic():
        DocumentPlacement.objects.filter(id=root.id).update(
            resolution_mode="pinned", pinned_revision_id=root.block.current_revision_id
        )
    with pytest.raises(DatabaseError), transaction.atomic():
        DocumentPlacement.objects.filter(id=root.id).delete()


@pytest.mark.django_db
def test_reuse_impact_shared_update_and_detach_preserve_live_and_local_semantics(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    source = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Shared standard", "markdown": "Revision one"},
        content_type="application/json",
    ).json()
    owner_client.post(
        reverse("msp-document-reference-list-create", kwargs={"document_entity_id": source["id"]}),
        {"organization_id": str(acme.entity_id)},
        content_type="application/json",
    )
    destination = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
        {"title": "Acme runbook", "markdown": "Introduction"},
        content_type="application/json",
    ).json()
    composed = owner_client.post(
        reverse(
            "organization-document-placement-list-create",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": destination["id"]},
        ),
        {"source_document_id": source["id"], "resolution_mode": "live"},
        content_type="application/json",
    ).json()
    placement = composed["placements"][1]
    reuse_url = reverse(
        "organization-document-placement-reuse",
        kwargs={
            "organization_entity_id": acme.entity_id,
            "document_entity_id": destination["id"],
            "placement_id": placement["id"],
        },
    )

    impact = owner_client.get(reuse_url)
    assert impact.status_code == 200
    assert impact.json()["can_edit_shared"] is True
    assert impact.json()["can_detach"] is True
    assert impact.json()["live_audience_count"] == 3
    assert {item["relationship"] for item in impact.json()["audiences"]} == {"source", "listing", "placement"}

    updated = owner_client.put(
        reuse_url,
        {"markdown": "Revision two", "base_revision_id": source["current_revision_id"]},
        content_type="application/json",
    )
    assert updated.status_code == 200
    assert updated.json()["resolved_markdown"].endswith("Revision two\n")
    assert BlockRevision.objects.filter(block__entity_id=source["block_id"]).count() == 2

    detached = owner_client.post(
        reverse(
            "organization-document-placement-detach",
            kwargs={
                "organization_entity_id": acme.entity_id,
                "document_entity_id": destination["id"],
                "placement_id": placement["id"],
            },
        )
    )
    assert detached.status_code == 200
    assert detached.json()["placements"][1]["block_id"] != source["block_id"]

    owner_client.put(
        reverse("msp-document-detail", kwargs={"document_entity_id": source["id"]}),
        {
            "title": source["title"],
            "markdown": "Revision three",
            "base_revision_id": updated.json()["placements"][1]["resolved_revision_id"],
        },
        content_type="application/json",
    )
    reloaded = owner_client.get(
        reverse(
            "organization-document-detail",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": destination["id"]},
        )
    ).json()
    assert reloaded["resolved_markdown"].endswith("Revision two\n")


@pytest.mark.django_db
def test_entity_mentions_search_and_render_are_workspace_scoped(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    beta = organization(installation.tenant, "Beta")
    acme_router = Entity.objects.create_owned(
        tenant=installation.tenant,
        organization=acme,
        entity_type="client_asset",
        display_name="Acme Core Router",
        visibility="client_visible",
    )
    beta_router = Entity.objects.create_owned(
        tenant=installation.tenant,
        organization=beta,
        entity_type="client_asset",
        display_name="Beta Core Router",
        visibility="client_visible",
    )
    search = owner_client.get(
        reverse("organization-document-mention-search", kwargs={"organization_entity_id": acme.entity_id}),
        {"q": "Core Router"},
    )
    assert search.status_code == 200
    assert [item["id"] for item in search.json()["results"]] == [str(acme_router.id)]

    markdown = (
        f"[authored](tekdocs://entity/{acme_router.id}) [secret authored label](tekdocs://entity/{beta_router.id})"
    )
    preview = owner_client.post(
        reverse("markdown-render"),
        {"markdown": markdown, "organization_id": str(acme.entity_id)},
        content_type="application/json",
    )
    assert preview.status_code == 200
    assert "Acme Core Router · client asset · Acme" in preview.json()["html"]
    assert "Unavailable reference" in preview.json()["html"]
    assert "secret authored label" not in preview.json()["html"]


@pytest.mark.django_db
def test_client_editor_cannot_change_msp_shared_block_but_can_detach(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    source = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "MSP standard", "markdown": "Canonical"},
        content_type="application/json",
    ).json()
    owner_client.post(
        reverse("msp-document-reference-list-create", kwargs={"document_entity_id": source["id"]}),
        {"organization_id": str(acme.entity_id)},
        content_type="application/json",
    )
    destination = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
        {"title": "Client guide", "markdown": "Local"},
        content_type="application/json",
    ).json()
    composed = owner_client.post(
        reverse(
            "organization-document-placement-list-create",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": destination["id"]},
        ),
        {"source_document_id": source["id"], "resolution_mode": "live"},
        content_type="application/json",
    ).json()
    placement = composed["placements"][1]

    editor = User.objects.create_user(email="client-editor@example.com", display_name="Client Editor")
    membership = TenantMembership.objects.create(
        tenant=installation.tenant,
        user=editor,
        role=BuiltInRole.READ_ONLY,
    )
    role = CustomRole.objects.create(
        tenant=installation.tenant,
        name="Client document editor",
        scope=CustomRoleScope.ORGANIZATION,
        created_by=installation.owner,
    )
    CustomRolePermission.objects.create(
        tenant=installation.tenant,
        role=role,
        permission="documents.edit",
    )
    ScopedRoleAssignment.objects.create(
        tenant=installation.tenant,
        membership=membership,
        role=role,
        organization=acme,
        created_by=installation.owner,
    )
    TOTP.activate(editor, generate_totp_secret())
    editor_client = Client()
    editor_client.force_login(editor)
    route_kwargs = {
        "organization_entity_id": acme.entity_id,
        "document_entity_id": destination["id"],
        "placement_id": placement["id"],
    }
    reuse_url = reverse("organization-document-placement-reuse", kwargs=route_kwargs)
    impact = editor_client.get(reuse_url)
    assert impact.status_code == 200
    assert impact.json()["can_edit_shared"] is False
    assert impact.json()["can_detach"] is True
    denied = editor_client.put(
        reuse_url,
        {"markdown": "Unauthorized", "base_revision_id": source["current_revision_id"]},
        content_type="application/json",
    )
    assert denied.status_code == 403
    assert BlockRevision.objects.filter(block__entity_id=source["block_id"]).count() == 1
    assert editor_client.post(reverse("organization-document-placement-detach", kwargs=route_kwargs)).status_code == 200

    local_block = editor_client.post(
        reverse(
            "organization-document-placement-list-create",
            kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": destination["id"]},
        ),
        {"operation": "create_block", "block_kind": "heading", "markdown": "## Client-owned"},
        content_type="application/json",
    )
    assert local_block.status_code == 200
    assert local_block.json()["placements"][-1]["block_kind"] == "heading"

    beta = organization(installation.tenant, "Beta")
    beta_document = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": beta.entity_id}),
        {"title": "Beta private", "markdown": "Primary"},
        content_type="application/json",
    ).json()
    cross_client = editor_client.post(
        reverse(
            "organization-document-placement-list-create",
            kwargs={"organization_entity_id": beta.entity_id, "document_entity_id": beta_document["id"]},
        ),
        {"operation": "create_block", "block_kind": "rich_text", "markdown": "Unauthorized"},
        content_type="application/json",
    )
    assert cross_client.status_code == 403


@pytest.mark.django_db
def test_publication_approval_role_is_exact_client_scoped_and_does_not_grant_withdrawal(
    owner_client, installation, tmp_path
):
    acme = organization(installation.tenant, "Approval Acme")
    beta = organization(installation.tenant, "Approval Beta")

    def publish_for(target):  # type: ignore[no-untyped-def]
        document = owner_client.post(
            reverse("organization-document-list-create", kwargs={"organization_entity_id": target.entity_id}),
            {"title": f"{target.entity.display_name} policy", "markdown": "Controlled"},
            content_type="application/json",
        ).json()
        publication = owner_client.post(
            reverse(
                "organization-document-publication-list-create",
                kwargs={"organization_entity_id": target.entity_id, "document_entity_id": document["id"]},
            ),
            {"reason": "Client distribution request", "audience": "client_visible", "retention": "permanent"},
            content_type="application/json",
        ).json()
        return document, publication

    with override_settings(MEDIA_ROOT=tmp_path):
        acme_document, acme_publication = publish_for(acme)
        beta_document, beta_publication = publish_for(beta)

    approver = User.objects.create_user(email="scoped-approver@example.invalid", display_name="Scoped Approver")
    membership = TenantMembership.objects.create(
        tenant=installation.tenant,
        user=approver,
        role=BuiltInRole.READ_ONLY,
    )
    role = CustomRole.objects.create(
        tenant=installation.tenant,
        name="Client publication approver",
        scope=CustomRoleScope.ORGANIZATION,
        created_by=installation.owner,
    )
    CustomRolePermission.objects.create(
        tenant=installation.tenant,
        role=role,
        permission="documents.approve",
    )
    ScopedRoleAssignment.objects.create(
        tenant=installation.tenant,
        membership=membership,
        role=role,
        organization=acme,
        created_by=installation.owner,
    )
    TOTP.activate(approver, generate_totp_secret())
    browser = Client()
    browser.force_login(approver)

    acme_kwargs = {
        "organization_entity_id": acme.entity_id,
        "document_entity_id": acme_document["id"],
        "publication_entity_id": acme_publication["id"],
    }
    approved = browser.post(
        reverse("organization-document-publication-approve", kwargs=acme_kwargs),
        {"reason": "Approved within assigned client scope"},
        content_type="application/json",
    )
    assert approved.status_code == 200
    assert approved.json()["lifecycle_state"] == "published"
    assert (
        browser.post(
            reverse("organization-document-publication-withdraw", kwargs=acme_kwargs),
            {"reason": "Not granted"},
            content_type="application/json",
        ).status_code
        == 403
    )

    beta_response = browser.post(
        reverse(
            "organization-document-publication-approve",
            kwargs={
                "organization_entity_id": beta.entity_id,
                "document_entity_id": beta_document["id"],
                "publication_entity_id": beta_publication["id"],
            },
        ),
        {"reason": "Cross-client attempt"},
        content_type="application/json",
    )
    assert beta_response.status_code == 403
    assert DocumentPublication.objects.get(entity_id=beta_publication["id"]).lifecycle_state == "pending_approval"


@pytest.mark.django_db
def test_live_document_exports_supported_sanitized_and_deterministic_formats(owner_client, installation):
    fidelity_markdown = (Path(__file__).parent / "fixtures" / "docx-fidelity.md").read_bytes()
    created = owner_client.post(
        reverse("msp-document-import"),
        {
            "title": "Router & switch guide",
            "file": SimpleUploadedFile(
                "guide.md",
                b"<script>alert(1)</script>\n\n" + fidelity_markdown,
                content_type="text/markdown",
            ),
            "category": "guide",
            "is_template": "false",
        },
    ).json()
    route = reverse("msp-document-export", kwargs={"document_entity_id": created["id"]})

    markdown = owner_client.get(route, {"export_format": "md"})
    html = owner_client.get(route, {"export_format": "html"})
    pdf = owner_client.get(route, {"export_format": "pdf"})
    docx = owner_client.get(route, {"export_format": "docx"})
    repeated_docx = owner_client.get(route, {"export_format": "docx"})

    assert markdown.status_code == 200
    assert markdown.content == created["resolved_markdown"].encode()
    assert b"<script>" not in html.content
    assert b"&lt;script&gt;alert(1)&lt;/script&gt;" in html.content
    assert pdf.content.startswith(b"%PDF-")
    assert docx.content.startswith(b"PK")
    assert docx.content == repeated_docx.content
    with ZipFile(BytesIO(docx.content)) as archive:
        document_xml = archive.read("word/document.xml")
        for xml_name in (name for name in archive.namelist() if name.endswith(".xml")):
            ElementTree.fromstring(archive.read(xml_name))
    reopened = WordDocument(BytesIO(docx.content))
    assert b"Router &amp; switch guide" in document_xml
    assert b'<w:pStyle w:val="Heading1"' in document_xml
    assert b"Connect uplink" in document_xml
    assert b"Preserve the rollback path" in document_xml
    assert b"ping gateway" in document_xml
    assert b"w:tbl" in document_xml
    assert b"w:hyperlink" in document_xml
    assert b"w:strike" in document_xml
    assert reopened.tables[0].cell(1, 1).text == "20"
    assert docx["Cache-Control"] == "private, no-store"
    assert docx["X-Content-Type-Options"] == "nosniff"
    assert docx["X-TekDocs-Export-Class"] == "editable_revision_snapshot"
    assert len(docx["X-TekDocs-Content-Digest"]) == 64
    metadata = list(AuditEvent.objects.filter(action="document.exported").values_list("metadata", flat=True))
    assert sorted(item["format"] for item in metadata) == ["docx", "docx", "html", "md", "pdf"]


@pytest.mark.django_db
def test_portable_bundle_freezes_exact_revisions_and_includes_only_selected_clean_files(
    owner_client, installation, tmp_path, monkeypatch
):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Portable network guide", "markdown": "# Network\n\nRetain this revision."},
        content_type="application/json",
    ).json()
    upload_route = reverse(
        "msp-document-attachment-list-create",
        kwargs={"document_entity_id": created["id"]},
    )
    export_route = reverse("msp-document-export", kwargs={"document_entity_id": created["id"]})

    with override_settings(MEDIA_ROOT=tmp_path):
        selected = owner_client.post(
            upload_route,
            {"file": SimpleUploadedFile("switch notes.txt", b"private attachment")},
        ).json()
        unselected = owner_client.post(
            upload_route,
            {"file": SimpleUploadedFile("diagram.txt", b"not exported")},
        ).json()
        query = {"export_format": "bundle", "attachment_ids": [selected["id"]]}
        first = owner_client.get(export_route, query)
        repeated = owner_client.get(export_route, query)
        monkeypatch.setattr("apps.core.document_exports.MAX_BUNDLE_ATTACHMENT_BYTES", 10)
        over_limit = owner_client.get(export_route, query)

    assert first.status_code == 200
    assert first.content == repeated.content
    assert over_limit.status_code == 409
    assert over_limit.content == b"Portable export file content may not exceed 50 MiB."
    assert first["Content-Type"] == "application/zip"
    assert "portable-network-guide-editable.zip" in first["Content-Disposition"]
    assert first["X-TekDocs-Export-Class"] == "editable_revision_snapshot"
    assert len(first["X-TekDocs-Content-Digest"]) == 64

    with ZipFile(BytesIO(first.content)) as archive:
        names = archive.namelist()
        assert names == sorted(names)
        manifest = json.loads(archive.read("manifest.json"))
        assert archive.read("document/document.md") == created["resolved_markdown"].encode()
        assert b"<h1>Network</h1>" in archive.read("document/document.html")
        attachment_path = manifest["attachments"][0]["path"]
        assert archive.read(attachment_path) == b"private attachment"
        assert all(unselected["id"] not in name for name in names)

    document = Document.objects.get(entity_id=created["id"])
    resolved_placement = document.placements.get()
    resolved_revision = resolved_placement.block.current_revision
    assert manifest["format"] == "tekdocs-portable-document/v3"
    assert manifest["export_class"] == "editable_revision_snapshot"
    assert manifest["immutable_publication"] is False
    assert manifest["content_digest"] == first["X-TekDocs-Content-Digest"]
    assert manifest["placements"] == [
        {
            "id": str(resolved_placement.id),
            "parent_id": None,
            "position": 0,
            "depth": 0,
            "block_id": str(resolved_placement.block.entity_id),
            "resolution_mode": "live",
            "audience_profile": "shared",
            "revision_id": str(resolved_revision.id),
            "revision_number": resolved_revision.revision_number,
            "checksum": resolved_revision.checksum,
        }
    ]
    assert manifest["attachments"][0]["id"] == selected["id"]
    assert manifest["attachments"][0]["checksum"] == sha256(b"private attachment").hexdigest()
    assert (
        AuditEvent.objects.filter(
            action="document.exported",
            metadata={"format": "bundle", "attachment_count": 1},
        ).count()
        == 2
    )


@pytest.mark.django_db(transaction=True)
def test_export_snapshot_serializes_with_concurrent_document_edit(installation, monkeypatch):
    client = Client()
    client.force_login(installation.owner)
    created = client.post(
        reverse("msp-document-list-create"),
        {"title": "Concurrent export", "markdown": "Revision one"},
        content_type="application/json",
    ).json()
    route = reverse("msp-document-export", kwargs={"document_entity_id": created["id"]})
    detail_route = reverse("msp-document-detail", kwargs={"document_entity_id": created["id"]})
    export_holds_lock = threading.Event()
    release_export = threading.Event()
    original_lock = document_exports.lock_document_composition

    def paused_lock(document):  # type: ignore[no-untyped-def]
        result = original_lock(document)
        export_holds_lock.set()
        assert release_export.wait(timeout=10)
        return result

    monkeypatch.setattr("apps.core.document_exports.lock_document_composition", paused_lock)
    results: dict[str, Any] = {}

    def run_export() -> None:
        close_old_connections()
        worker = Client()
        worker.force_login(installation.owner)
        results["export"] = worker.get(route, {"export_format": "md"})
        close_old_connections()

    def run_edit() -> None:
        close_old_connections()
        worker = Client()
        worker.force_login(installation.owner)
        results["edit"] = worker.put(
            detail_route,
            {
                "title": "Concurrent export",
                "markdown": "Revision two",
                "base_revision_id": created["current_revision_id"],
                "category": "general",
            },
            content_type="application/json",
        )
        close_old_connections()

    export_thread = threading.Thread(target=run_export)
    edit_thread = threading.Thread(target=run_edit)
    export_thread.start()
    assert export_holds_lock.wait(timeout=10)
    edit_thread.start()
    time.sleep(0.2)
    assert edit_thread.is_alive()
    release_export.set()
    export_thread.join(timeout=10)
    edit_thread.join(timeout=10)

    assert not export_thread.is_alive()
    assert not edit_thread.is_alive()
    assert results["export"].status_code == 200
    assert results["export"].content == b"Revision one\n"
    assert results["edit"].status_code == 200
    monkeypatch.setattr("apps.core.document_exports.lock_document_composition", original_lock)
    assert client.get(route, {"export_format": "md"}).content == b"Revision two\n"


@pytest.mark.django_db
def test_file_download_and_bundle_fail_closed_during_storage_outage(owner_client, installation, tmp_path):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Storage outage", "markdown": "Retained source"},
        content_type="application/json",
    ).json()
    upload_route = reverse("msp-document-attachment-list-create", kwargs={"document_entity_id": created["id"]})
    with override_settings(MEDIA_ROOT=tmp_path):
        attachment = owner_client.post(
            upload_route,
            {"file": SimpleUploadedFile("outage.txt", b"private retained bytes")},
        ).json()
        with override_settings(
            TEKDOCS_ATTACHMENT_STORAGE_PROVIDER="apps.core.tests.test_documents.UnavailableStorageProvider"
        ):
            download = owner_client.get(
                reverse(
                    "msp-document-attachment-download",
                    kwargs={"document_entity_id": created["id"], "attachment_entity_id": attachment["id"]},
                )
            )
            bundle = owner_client.get(
                reverse("msp-document-export", kwargs={"document_entity_id": created["id"]}),
                {"export_format": "bundle", "attachment_ids": [attachment["id"]]},
            )

    assert download.status_code == 400
    assert bundle.status_code == 409
    assert bundle.content == b"A selected file failed its retained-content integrity check."
    assert not AuditEvent.objects.filter(action__in=("document.exported", "document.attachment.downloaded")).exists()


@pytest.mark.django_db
def test_portable_bundle_rejects_foreign_or_unavailable_files_without_disclosing_them(
    owner_client, installation, tmp_path
):
    acme = organization(installation.tenant, "Acme portable export")
    beta = organization(installation.tenant, "Beta portable export")
    first_document = owner_client.post(
        reverse(
            "organization-document-list-create",
            kwargs={"organization_entity_id": acme.entity_id},
        ),
        {"title": "First document", "markdown": "First"},
        content_type="application/json",
    ).json()
    second_document = owner_client.post(
        reverse(
            "organization-document-list-create",
            kwargs={"organization_entity_id": beta.entity_id},
        ),
        {"title": "Second document", "markdown": "Second"},
        content_type="application/json",
    ).json()
    with override_settings(MEDIA_ROOT=tmp_path):
        foreign_file = owner_client.post(
            reverse(
                "organization-document-attachment-list-create",
                kwargs={
                    "organization_entity_id": beta.entity_id,
                    "document_entity_id": second_document["id"],
                },
            ),
            {"file": SimpleUploadedFile("foreign.txt", b"foreign private content")},
        ).json()
        rejected = owner_client.get(
            reverse(
                "organization-document-export",
                kwargs={
                    "organization_entity_id": acme.entity_id,
                    "document_entity_id": first_document["id"],
                },
            ),
            {"export_format": "bundle", "attachment_ids": [foreign_file["id"]]},
        )

    assert rejected.status_code == 409
    assert rejected.content == b"One or more selected or referenced files are unavailable."
    assert not AuditEvent.objects.filter(action="document.exported", metadata__format="bundle").exists()


@pytest.mark.django_db
def test_static_exports_use_retained_snapshot_and_pdf_artifact(owner_client, installation, tmp_path):
    created = owner_client.post(
        reverse("msp-document-import"),
        {
            "title": "Frozen guide",
            "file": SimpleUploadedFile("guide.md", b"# Original\n\nRetained content"),
            "category": "guide",
            "is_template": "false",
        },
    ).json()
    publication_route = reverse("msp-document-publication-list-create", kwargs={"document_entity_id": created["id"]})
    with override_settings(MEDIA_ROOT=tmp_path):
        publication = owner_client.post(
            publication_route,
            {"reason": "Release baseline", "audience": "msp_internal", "retention": "permanent"},
            content_type="application/json",
        ).json()
        owner_client.put(
            reverse("msp-document-detail", kwargs={"document_entity_id": created["id"]}),
            {
                "title": "Changed guide",
                "markdown": "Replacement content",
                "base_revision_id": created["current_revision_id"],
                "category": "guide",
            },
            content_type="application/json",
        )
        route = reverse(
            "msp-document-publication-export",
            kwargs={"document_entity_id": created["id"], "publication_entity_id": publication["id"]},
        )
        markdown = owner_client.get(route, {"export_format": "md"})
        html = owner_client.get(route, {"export_format": "html"})
        pdf = owner_client.get(route, {"export_format": "pdf"})
        docx = owner_client.get(route, {"export_format": "docx"})
        bundle = owner_client.get(route, {"export_format": "bundle"})
        retained = DocumentPublication.objects.get(entity_id=publication["id"])
        pdf_artifact = retained.artifacts.get(kind="pdf")

        assert markdown.content == retained.canonical_markdown.encode()
        assert "frozen-guide-static.md" in markdown["Content-Disposition"]
        assert markdown["X-TekDocs-Export-Class"] == "immutable_static_publication"
        assert markdown["X-TekDocs-Content-Digest"] == retained.content_digest
        assert retained.sanitized_html.encode() in html.content
        assert html["X-TekDocs-Export-Class"] == "immutable_static_publication"
        assert pdf.content == pdf_artifact.file.read()
        assert pdf["X-TekDocs-Content-Digest"] == retained.content_digest
        with ZipFile(BytesIO(docx.content)) as archive:
            document_xml = archive.read("word/document.xml")
        assert b"Retained content" in document_xml
        assert b"Replacement content" not in document_xml
        assert docx["X-TekDocs-Export-Class"] == "immutable_static_publication"
        assert bundle.status_code == 400
        assert AuditEvent.objects.filter(action="document.publication_exported", metadata={"format": "docx"}).exists()


@pytest.mark.django_db
def test_document_export_rejects_unknown_format(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-import"),
        {"title": "Guide", "file": SimpleUploadedFile("guide.md", b"Body"), "is_template": "false"},
    ).json()
    route = reverse("msp-document-export", kwargs={"document_entity_id": created["id"]})

    response = owner_client.get(route, {"export_format": "gdoc"})

    assert response.status_code == 400


@pytest.mark.django_db
def test_legacy_document_restructure_is_previewed_atomic_audited_and_idempotent(owner_client, installation):
    original = "# Network standard\n\nUse segmented networks.\n\n## Required VLANs\n\n- Staff\n- Guest\n"
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Network standard", "markdown": original, "category": "guide"},
        content_type="application/json",
    ).json()
    document = Document.objects.get(entity_id=created["id"])
    original_entity_id = document.entity_id
    original_revision = BlockRevision.objects.get(id=created["current_revision_id"])
    route = reverse("msp-document-restructure", kwargs={"document_entity_id": created["id"]})

    before = {
        "placements": DocumentPlacement.objects.filter(document=document).count(),
        "revisions": BlockRevision.objects.filter(block__source_document=document).count(),
        "audits": AuditEvent.objects.filter(action="document.semantic_sections_created").count(),
    }
    preview = owner_client.get(route)

    assert preview.status_code == 200
    payload = preview.json()
    assert payload["eligible"] is True
    assert payload["base_revision_id"] == created["current_revision_id"]
    assert payload["section_count"] == 4
    assert [section["kind"] for section in payload["sections"]] == [
        "heading",
        "rich_text",
        "heading",
        "rich_text",
    ]
    assert before == {
        "placements": DocumentPlacement.objects.filter(document=document).count(),
        "revisions": BlockRevision.objects.filter(block__source_document=document).count(),
        "audits": AuditEvent.objects.filter(action="document.semantic_sections_created").count(),
    }

    converted = owner_client.post(
        route,
        {"base_revision_id": created["current_revision_id"]},
        content_type="application/json",
    )
    assert converted.status_code == 200
    result = converted.json()
    assert result["status"] == "restructured"
    assert result["section_count"] == 4
    assert result["document"]["id"] == str(original_entity_id)
    assert result["document"]["category"] == "guide"
    assert result["document"]["placement_count"] == 4
    assert (
        result["document"]["resolved_markdown"]
        == "\n\n".join(section["markdown"] for section in payload["sections"]) + "\n"
    )
    original_revision.refresh_from_db()
    primary_revisions = list(BlockRevision.objects.filter(block=original_revision.block).order_by("revision_number"))
    assert [revision.id for revision in primary_revisions] == [
        original_revision.id,
        uuid.UUID(result["document"]["current_revision_id"]),
    ]
    assert primary_revisions[1].parent_id == original_revision.id
    audit = AuditEvent.objects.get(action="document.semantic_sections_created", entity_id=original_entity_id)
    assert audit.metadata["base_revision_id"] == str(original_revision.id)
    assert audit.metadata["section_count"] == 4
    assert original not in str(audit.metadata)

    replay = owner_client.post(
        route,
        {"base_revision_id": created["current_revision_id"]},
        content_type="application/json",
    )
    assert replay.status_code == 200
    assert replay.json()["status"] == "already_restructured"
    assert DocumentPlacement.objects.filter(document=document).count() == 4
    assert (
        AuditEvent.objects.filter(action="document.semantic_sections_created", entity_id=original_entity_id).count()
        == 1
    )


@pytest.mark.django_db
def test_legacy_document_restructure_rejects_stale_shared_and_managed_content(owner_client, installation):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Legacy guide", "markdown": "# Start\n\nFirst.\n\n## Next\n\nSecond."},
        content_type="application/json",
    ).json()
    route = reverse("msp-document-restructure", kwargs={"document_entity_id": created["id"]})
    updated = owner_client.put(
        reverse("msp-document-detail", kwargs={"document_entity_id": created["id"]}),
        {
            "title": "Legacy guide",
            "markdown": "# Start\n\nChanged.\n\n## Next\n\nSecond.",
            "base_revision_id": created["current_revision_id"],
        },
        content_type="application/json",
    )
    assert updated.status_code == 200
    stale = owner_client.post(
        route,
        {"base_revision_id": created["current_revision_id"]},
        content_type="application/json",
    )
    assert stale.status_code == 409
    assert stale.json()["code"] == "revision_conflict"
    assert Document.objects.get(entity_id=created["id"]).placements.count() == 1

    destination = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Destination", "markdown": "Local"},
        content_type="application/json",
    ).json()
    reused = owner_client.post(
        reverse("msp-document-placement-list-create", kwargs={"document_entity_id": destination["id"]}),
        {"source_document_id": created["id"], "resolution_mode": "live"},
        content_type="application/json",
    )
    assert reused.status_code == 200
    shared_preview = owner_client.get(route).json()
    assert shared_preview["eligible"] is False
    assert "shared_primary" in {item["code"] for item in shared_preview["blockers"]}

    managed = owner_client.post(
        reverse("msp-document-list-create"),
        {
            "title": "Managed template",
            "markdown": "# Start\n\nFirst.\n\n## Next\n\nSecond.",
            "is_template": True,
        },
        content_type="application/json",
    ).json()
    managed_preview = owner_client.get(
        reverse("msp-document-restructure", kwargs={"document_entity_id": managed["id"]})
    ).json()
    assert managed_preview["eligible"] is False
    assert "template_source" in {item["code"] for item in managed_preview["blockers"]}


@pytest.mark.django_db(transaction=True)
def test_legacy_document_restructure_rolls_back_without_orphans(owner_client, installation, monkeypatch):
    created = owner_client.post(
        reverse("msp-document-list-create"),
        {"title": "Rollback guide", "markdown": "# Start\n\nFirst.\n\n## Next\n\nSecond."},
        content_type="application/json",
    ).json()
    document = Document.objects.get(entity_id=created["id"])
    original_create = DocumentPlacement.objects.create

    def fail_new_placement(**kwargs):  # type: ignore[no-untyped-def]
        if kwargs.get("position") == 1:
            raise RuntimeError("injected placement failure")
        return original_create(**kwargs)

    monkeypatch.setattr(DocumentPlacement.objects, "create", fail_new_placement)
    with pytest.raises(RuntimeError, match="injected placement failure"):
        restructure_document(
            document=document,
            actor_id=installation.owner.id,
            base_revision_id=uuid.UUID(created["current_revision_id"]),
        )

    document.refresh_from_db()
    primary = document.placements.select_related("block").get(parent__isnull=True, position=0)
    assert document.placements.count() == 1
    assert primary.block.current_revision_id == uuid.UUID(created["current_revision_id"])
    assert Block.objects.filter(source_document=document).count() == 1
    assert BlockRevision.objects.filter(block__source_document=document).count() == 1
    assert not AuditEvent.objects.filter(
        action="document.semantic_sections_created", entity_id=document.entity_id
    ).exists()


@pytest.mark.django_db
def test_client_document_restructure_respects_scoped_authorization(owner_client, installation):
    acme = organization(installation.tenant, "Acme")
    beta = organization(installation.tenant, "Beta")
    acme_document = owner_client.post(
        reverse("organization-document-list-create", kwargs={"organization_entity_id": acme.entity_id}),
        {"title": "Acme guide", "markdown": "# Start\n\nFirst.\n\n## Next\n\nSecond."},
        content_type="application/json",
    ).json()
    editor = User.objects.create_user(email="restructure-editor@example.com", display_name="Acme editor")
    membership = TenantMembership.objects.create(
        tenant=installation.tenant,
        user=editor,
        role=BuiltInRole.READ_ONLY,
    )
    role = CustomRole.objects.create(
        tenant=installation.tenant,
        name="Document restructurer",
        scope=CustomRoleScope.ORGANIZATION,
        created_by=installation.owner,
    )
    CustomRolePermission.objects.create(tenant=installation.tenant, role=role, permission="documents.edit")
    ScopedRoleAssignment.objects.create(
        tenant=installation.tenant,
        membership=membership,
        role=role,
        organization=acme,
        created_by=installation.owner,
    )
    TOTP.activate(editor, generate_totp_secret())
    editor_client = Client()
    editor_client.force_login(editor)
    acme_route = reverse(
        "organization-document-restructure",
        kwargs={"organization_entity_id": acme.entity_id, "document_entity_id": acme_document["id"]},
    )
    assert editor_client.get(acme_route).status_code == 200
    cross_client_route = reverse(
        "organization-document-restructure",
        kwargs={"organization_entity_id": beta.entity_id, "document_entity_id": acme_document["id"]},
    )
    assert editor_client.get(cross_client_route).status_code == 403
